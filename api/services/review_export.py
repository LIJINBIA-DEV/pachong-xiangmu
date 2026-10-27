"""Export requested post metadata while preserving source and review data."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from api.services.data_browser import load_review_dataset, text, web_url
from api.services.manual_review import annotations, note_key, image_download_problem
from weibo_metadata import metadata_platform, metadata_view

HEADERS = ['编号', '平台', '原帖链接', '标题', '发布时间', '地点（定位优先）', '命中关键词', '发布内容', '与关键词一相关的评论',
           '账号昵称', '账号 ID', '账号主页链接', '认证说明', '性别', 'IP 属地', '地点来源']
TIMEZONE = ZoneInfo('Asia/Shanghai')
PLATFORMS = {'wb': '微博', 'weibo': '微博', 'xhs': '小红书', 'dy': '抖音', 'douyin': '抖音',
    'ks': '快手', 'kuaishou': '快手', 'bili': '哔哩哔哩', 'bilibili': '哔哩哔哩', 'tieba': '百度贴吧', 'zhihu': '知乎'}


def _platform(note, path):
    saved = text(note.get('platform')).strip().lower()
    if saved in PLATFORMS:
        return PLATFORMS[saved]
    source = Path(note.get('_collection_file') or path)
    for part in source.parts:
        if part.lower() in PLATFORMS:
            return PLATFORMS[part.lower()]
    if source.name.lower().startswith('weibo_'):
        return '微博'
    return text(note.get('platform')) or '未注明'


def _published(note):
    for key in ('create_date_time', 'published_at', 'time'):
        if note.get(key) is not None and text(note[key]).strip():
            value = text(note[key])
            try:
                moment = datetime.fromisoformat(value)
                if moment.tzinfo is not None:
                    moment = moment.astimezone(TIMEZONE)
                return moment.strftime('%Y-%m-%d %H:%M:%S')
            except ValueError:
                return value
    value = note.get('create_time')
    if value not in (None, ''):
        try:
            seconds = float(value)
            if seconds > 100_000_000_000:
                seconds /= 1000
            return datetime.fromtimestamp(seconds, TIMEZONE).strftime('%Y-%m-%d %H:%M:%S')
        except (TypeError, ValueError, OverflowError, OSError):
            pass
    return '未采集发布时间'


def _location(note):
    # Network/IP location and places mentioned in prose are not posting locations.
    for key in ('publish_location', 'location', 'location_name', 'geo_location'):
        value = note.get(key)
        if isinstance(value, dict):
            value = value.get('name') or value.get('address')
        if isinstance(value, str) and value.strip():
            return value.strip()
    return '未采集发布地点'


def _post_comments(note, comments):
    rows = list(comments)
    if isinstance(note.get('matched_comment'), dict):
        rows.append(note['matched_comment'])
    nid = text(note.get('note_id') or note.get('aweme_id'))
    result, seen = [], set()
    for row in rows:
        if row.get('note_id') and text(row['note_id']) != nid:
            continue
        content = text(row.get('content')).strip()
        signature = (text(row.get('comment_id')), content)
        if content and signature not in seen:
            seen.add(signature)
            result.append(row)
    return result


def _comment_gap(note):
    if note.get('comments_collection_stop_reason') == 'comments_fetch_failed':
        return '评论抓取失败或接口不可用；未保存评论'
    if note.get('comments_collection_complete'):
        return '已读取评论接口，未取得可保存的评论'
    stop = note.get('semantic_comment_stop_reason')
    if stop in {'body_match', 'review_candidate_accepted'} and not note.get('comments_collected_count'):
        return '本轮正文初筛通过后跳过了评论抓取；未保存评论'
    try:
        count = int(note.get('comments_count'))
    except (TypeError, ValueError):
        count = None
    if count == 0:
        return '原帖显示0条评论；未保存评论'
    return '未采集到可供导出的评论' + (f'；原帖显示{count}条评论' if count is not None else '')


def _related_comments(note, comments):
    first = text(note.get('source_keyword')).strip().split(maxsplit=1)
    if not first:
        return '关键词一未注明，无法筛选'
    rows = _post_comments(note, comments)
    related = []
    for row in rows:
        content = text(row.get('content')).strip()
        if first[0].casefold() in content.casefold():
            author = text(row.get('nickname')).strip()
            related.append((f'【{author}】' if author else '') + content)
    if related:
        return '\n\n'.join(related)
    return f'已采集{len(rows)}条评论，未发现包含关键词一的评论；见评论明细' if rows else _comment_gap(note)


def export_review(root, path, *, snapshot=None):
    if snapshot is None:
        notes, comments = load_review_dataset(root, path)
        review = annotations(root, path, persist=False, dataset=(notes, comments))
    else:
        notes, comments, review = snapshot
    retained = [note for note in notes if review['entries'][note_key(note)]['decision'] == 'qualified']
    if not retained:
        raise ValueError('review_export_empty')
    book = Workbook()
    sheet = book.active
    sheet.title = '审核结果'
    sheet.append(HEADERS)
    detail = book.create_sheet('评论明细')
    detail.append(['审核结果编号', '帖子编号', '命中关键词', '评论编号', '评论人', '评论内容', '发布时间', '是否包含关键词一', '状态'])
    image_count = 0
    for index, note in enumerate(retained, 1):
        entry = review['entries'][note_key(note)]
        image_count += sum(not image_download_problem(root, img) for img in entry['review_images'])
        nid = text(note.get('note_id') or note.get('aweme_id'))
        metadata = metadata_view(note, metadata_platform(note, path))
        sheet.append([index, _platform(note, path), web_url(note.get('note_url') or note.get('url')) or '未采集原帖链接',
            text(note.get('title') or note.get('note_title')).strip() or '未采集独立标题',
            _published(note), metadata['display_location'], text(note.get('source_keyword')),
            text(note.get('content') or note.get('desc')), _related_comments(note, comments.get(nid, [])),
            text(note.get('nickname')) or '未获取', metadata['account_id'] or '未获取',
            metadata['profile_url'] or '未获取', metadata['verification'], metadata['gender'],
            metadata['post_ip_region'] or '未获取', metadata['location_source']])
        # Keep all crawled text literal, including strings beginning with '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[sheet.max_row].height = 84
        saved_comments = _post_comments(note, comments.get(nid, []))
        first = text(note.get('source_keyword')).strip().split(maxsplit=1)
        if not saved_comments:
            detail.append([index, nid, text(note.get('source_keyword')), '', '', '', '', '', _comment_gap(note)])
        for comment in saved_comments:
            content = text(comment.get('content'))
            related = bool(first and first[0].casefold() in content.casefold())
            # A cell cannot hold more than 32767 characters; keep the full comment in numbered parts.
            parts = [content[start:start + 32000] for start in range(0, len(content), 32000)] or ['']
            for part, fragment in enumerate(parts, 1):
                detail.append([index, nid, text(note.get('source_keyword')), text(comment.get('comment_id')),
                    text(comment.get('nickname')), fragment, _published(comment), '是' if related else '否',
                    '已采集' if len(parts) == 1 else f'已采集，分段{part}/{len(parts)}'])
    for row in detail:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    for column, width in enumerate([18, 26, 22, 26, 24, 90, 24, 20, 50], 1):
        detail.column_dimensions[get_column_letter(column)].width = width
    for cell in detail[1]:
        cell.fill = PatternFill('solid', fgColor='0F172A')
        cell.font = Font(color='FFFFFF', bold=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='0F172A')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for index, width in enumerate([10, 12, 45, 28, 24, 30, 22, 70, 70, 24, 25, 45, 35, 12, 20, 18], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    book.close()
    output.seek(0)
    filename = path.stem + '_审核结果_' + datetime.now(TIMEZONE).strftime('%Y%m%d_%H%M%S') + '.xlsx'
    return output, filename, {'retained': len(retained), 'images': image_count, 'pending': review['total'] - review['completed']}
