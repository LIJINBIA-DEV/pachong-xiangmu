"""Prepare disk-backed ZIP64 delivery packages without changing annotations."""
from datetime import datetime
import hashlib
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_STORED

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from api.services.data_browser import load_review_dataset, text
from api.services.manual_review import annotations, image_download_problem, note_key
from api.services.review_export import export_review, _platform, TIMEZONE
from image_subject import local_image_path

MAPPING_HEADERS = ['审核结果编号', '平台', '帖子 ID', '原图片名', '包内相对路径', '导出状态']


def safe_component(value):
    original = text(value)
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', original).strip(' .') or '未注明'
    if re.fullmatch(r'(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', cleaned, re.I):
        cleaned = '_' + cleaned
    if cleaned != original or len(cleaned.encode('utf-8')) > 180:
        suffix = Path(cleaned).suffix
        if len(suffix.encode('utf-8')) > 20:
            suffix = ''
        stem = cleaned[:-len(suffix)] if suffix else cleaned
        stem = stem.encode('utf-8')[:140].decode('utf-8', errors='ignore')
        cleaned = stem + '_' + hashlib.sha256(original.encode()).hexdigest()[:10] + suffix
    return cleaned


def build_package(root, path, output, cancelled, progress):
    notes, comments = load_review_dataset(root, path)
    review = annotations(root, path, persist=False, dataset=(notes, comments))
    retained = [n for n in notes if review['entries'][note_key(n)]['decision'] == 'qualified']
    if not retained:
        raise ValueError('review_export_empty')
    excel, _, counts = export_review(root, path, snapshot=(notes, comments, review))
    book = load_workbook(excel)
    mapping = book.create_sheet('图片对应表')
    mapping.append(MAPPING_HEADERS)
    counts.update(images=0, missing=0, no_images=0)
    selected = sum(sum(e['images'].get(i['key']) == 'qualified' for i in e['review_images'])
                   for n in retained for e in [review['entries'][note_key(n)]])
    progress(0, selected)
    used, packed = {}, {}
    processed = 0
    try:
        with ZipFile(output, 'w', compression=ZIP_STORED, allowZip64=True) as archive:
            for index, note in enumerate(retained, 1):
                if cancelled.is_set():
                    raise ValueError('review_export_cancelled')
                entry = review['entries'][note_key(note)]
                approved = [i for i in entry['review_images'] if entry['images'].get(i['key']) == 'qualified']
                platform, nid = _platform(note, path), text(note.get('note_id') or note.get('aweme_id'))
                success = 0
                for image in approved:
                    if cancelled.is_set():
                        raise ValueError('review_export_cancelled')
                    relative, status = '', '已导出'
                    source = local_image_path(root, image)
                    problem = image_download_problem(root, image)
                    if problem or source is None:
                        status = '未下载' if problem == 'not_downloaded' else '图片损坏或下载失败'
                    else:
                        # Resolve again immediately before opening. Do not follow links out of the data root.
                        source = source.resolve()
                        source.relative_to(root.resolve())
                        signature = (str(source), source.stat().st_mtime_ns, source.stat().st_size)
                        if signature in packed:
                            relative = packed[signature]
                        else:
                            candidate = '/'.join(['图片', safe_component(platform), safe_component(nid), safe_component(image['name'])])
                            if candidate.casefold() in used:
                                item = Path(candidate)
                                candidate = str(item.with_name(item.stem + '_' + hashlib.sha256(str(source).encode()).hexdigest()[:12] + item.suffix)).replace('\\', '/')
                            # Opening failure can be reported as a missing image; write failure must abort the ZIP.
                            try:
                                handle = source.open('rb')
                            except OSError:
                                status = '原图无法读取'
                            else:
                                with handle, archive.open(candidate, 'w', force_zip64=True) as target:
                                    while chunk := handle.read(1024 * 1024):
                                        if cancelled.is_set():
                                            raise ValueError('review_export_cancelled')
                                        target.write(chunk)
                                after = source.stat()
                                if (after.st_mtime_ns, after.st_size) != signature[1:]:
                                    raise ValueError('review_export_source_changed')
                                used[candidate.casefold()] = signature
                                packed[signature] = candidate
                                relative = candidate
                                counts['images'] += 1
                    if not relative:
                        counts['missing'] += 1
                    else:
                        success += 1
                    mapping.append([index, platform, nid, image['name'], relative, status])
                    processed += 1
                    progress(processed, selected)
                if not success:
                    counts['no_images'] += 1
                    if not approved:
                        mapping.append([index, platform, nid, '', '', '无可导出合格图片'])
            mapping.freeze_panes = 'A2'
            mapping.auto_filter.ref = mapping.dimensions
            for row in mapping:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            for col, width in zip('ABCDEF', [18, 14, 25, 32, 70, 28]):
                mapping.column_dimensions[col].width = width
            workbook_path = output.with_suffix('.xlsx')
            try:
                book.save(workbook_path)
                archive.write(workbook_path, '审核结果.xlsx')
            finally:
                workbook_path.unlink(missing_ok=True)
            description = '\n'.join([
                '审核结果导出说明', '来源任务：' + path.name,
                '导出时间：' + datetime.now(TIMEZONE).strftime('%Y-%m-%d %H:%M:%S'),
                '结果范围：' + ('部分结果' if counts['pending'] or counts['missing'] else '当前已保留结果'),
                f"保留帖子：{counts['retained']}", f"已打包原图：{counts['images']}",
                f"待审核帖子：{counts['pending']}", f"合格图片缺失：{counts['missing']}",
                f"无可导出合格图片的保留帖子：{counts['no_images']}",
                '具体图片路径和缺失原因请查看 Excel 的“图片对应表”。',
                '重复引用共用包内原图；不合格或未标记图片不属于缺失图片。',
                '导出不修改原帖、审核标记及识别缓存，不重新识别或补下载。',
            ])
            archive.writestr('导出说明.txt', description.encode('utf-8-sig'))
        return counts
    finally:
        book.close()

