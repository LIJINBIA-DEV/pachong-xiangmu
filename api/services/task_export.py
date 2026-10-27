"""Export one explicitly selected post status, keeping source files untouched."""
from io import BytesIO
import json

from openpyxl import Workbook
from api.services.data_browser import load_review_dataset, post_id, reviewed_notes, collection_status


def export_task_posts(root, path, status):
    notes, comments = load_review_dataset(root, path)
    notes = [note for note in reviewed_notes(root, path, notes, comments) if collection_status(note) == status]
    ids = {post_id(note) for note in notes}
    unique_comments = {}
    for nid in ids:
        for comment in comments.get(nid, []):
            key = (nid, str(comment.get('comment_id', '')), str(comment.get('content', '')))
            unique_comments.setdefault(key, comment)
    book = Workbook()
    book.remove(book.active)
    for title, rows in [('帖子', notes), ('评论', list(unique_comments.values()))]:
        sheet = book.create_sheet(title)
        columns = list(dict.fromkeys(key for row in rows for key in row if not key.startswith('_')))
        sheet.append(columns or ['note_id'])
        for row in rows:
            values = [json.dumps(row.get(key), ensure_ascii=False) if isinstance(row.get(key), (dict, list)) else row.get(key) for key in columns]
            sheet.append(values)
            # Collected text must remain text even when it starts with '='.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
