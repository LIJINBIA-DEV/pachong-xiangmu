"""Run inside the deployed container; synthetic checks do not collect real posts."""
import asyncio
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


async def main():
    from api.main import app
    from api.services.manual_review import update_annotation, annotations
    from api.services.review_export import export_review, HEADERS
    from image_subject import image_suggestion, valid_prediction
    from image_subject_cloud import SiliconFlowImageReviewer
    from image_subject_settings import runtime_settings
    import config

    root = Path('/app/data')
    assert root.is_dir() and os.access(root, os.W_OK)
    assert config.LOGIN_TYPE == 'cookie' and config.HEADLESS and not config.ENABLE_CDP_MODE
    options = runtime_settings(root)
    assert options['provider'] == 'siliconflow' and options['api_key']
    with TestClient(app) as client:
        assert client.get('/api/health').json()['status'] == 'ok'
        assert client.get('/api/crawler/status').json()['status'] == 'idle'
        settings = client.get('/api/data/image-subject/settings').json()
        assert settings['provider'] == 'siliconflow' and settings['ready']
        assert 'api_key' not in settings
        assert client.post('/api/crawler/start', json={'platform': 'wb', 'cookies': '',
            'image_subject_recognition': False}).status_code == 400

    with tempfile.TemporaryDirectory(prefix='.deployment-check-', dir=root) as folder:
        test_root = Path(folder)
        image_path = test_root / 'poster.png'
        image = Image.new('RGB', (480, 480), '#dfefff')
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 42)
        ImageDraw.Draw(image).multiline_text((25, 105), 'A HAPPY DAY\n\nENJOY LIFE\n\nBE YOURSELF', font=font, fill='black', spacing=20)
        image.save(image_path)
        reviewer = SiliconFlowImageReviewer(test_root, options['api_key'])
        started = time.perf_counter()
        try:
            await reviewer.start()
            result = await reviewer.classify(image_path)
        finally:
            await reviewer.close()
        assert valid_prediction(result) and image_suggestion(result) == 'unqualified'

        path = test_root / 'weibo' / 'checks.json'
        path.parent.mkdir()
        path.write_text(json.dumps([{'note_id':'deployment-check','creator_hash':'deployment-check-user',
            'source_keyword':'可爱 公园','content':'部署验证：这个公园的花箱非常可爱。','image_names':'001.jpg',
            'semantic_result':'matched'}],ensure_ascii=False),encoding='utf-8')
        media = test_root / 'wb/media/deployment-check'
        media.mkdir(parents=True)
        Image.new('RGB',(80,80),'green').save(media/'001.jpg')
        key = 'deployment-check:可爱 公园'
        update_annotation(test_root,path,key,'account','non_marketing')
        update_annotation(test_root,path,key,'image','qualified','001.jpg')
        before = path.read_bytes()
        exported,_,counts = export_review(test_root,path)
        book = load_workbook(BytesIO(exported.getvalue()))
        assert [cell.value for cell in book.active[1]] == HEADERS
        assert book.active.max_row == 2 and counts['retained'] == 1
        book.close()
        assert path.read_bytes() == before
        assert annotations(test_root,path)['entries'][key]['images']['001.jpg'] == 'qualified'
        print(json.dumps({'cloud_model':options['model'],'provider':result['provider'],
            'sample_label':result['label'],'sample_type':result['image_type'],
            'sample_suggestion':image_suggestion(result),'sample_seconds':round(time.perf_counter()-started,2),
            'excel_export':'passed','manual_marks':'passed','data_volume_write':'passed',
            'cookie_required':'passed'},ensure_ascii=False))


if __name__ == '__main__':
    asyncio.run(main())
