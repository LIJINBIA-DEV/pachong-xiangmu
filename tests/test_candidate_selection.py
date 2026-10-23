"""Replay seven user-selected scene candidates; no live model accuracy claim."""
import json
from pathlib import Path

import httpx
import openpyxl
import pytest
from types import SimpleNamespace
from unittest.mock import AsyncMock

from semantic_filter import JevReviewer, StreamingSemanticFilter
from tests.semantic_responses import grounded_response
from api.services.data_browser import browse_posts

SAMPLES = json.loads((Path(__file__).parent / 'fixtures/weibo_candidates_20261008.json').read_text(encoding='utf-8'))


def replay(sample):
    def respond(request):
        payload = json.loads(request.content)
        if 'support' in payload['questions']:
            return httpx.Response(200, json={'answers': {'support': {'choice': 'unclear', 'confidence': sample['verification_confidence']}}})
        # Old low confidence and inconclusive verification are deliberately
        # preserved, rather than replacing them with a fabricated strong vote.
        choice = 'keep' if sample['quality_reason'] else 'review'
        result = grounded_response(request, choice, confidence=sample['semantic_confidence'],
                                   kind=sample['semantic_object_type'], relation=sample['semantic_target_relation'])
        answers = result.json()
        options = payload['questions']['evidence_text']['criteria']
        for field, saved in [('evidence_text', 'semantic_evidence'), ('target_evidence', 'target_evidence')]:
            exact = next((key for key, value in options.items() if value.endswith(' ' + sample[saved])), None)
            assert exact, f"Original evidence missing for {sample['note_id']} {field}"
            answers['answers'][field]['choice'] = exact
        return httpx.Response(200, json=answers)
    return respond


@pytest.mark.asyncio
@pytest.mark.parametrize('sample', SAMPLES, ids=[s['note_id'] for s in SAMPLES])
async def test_user_examples_pass_initial_screen_without_comments_and_preserve_old_scores(tmp_path, sample):
    async with httpx.AsyncClient(transport=httpx.MockTransport(replay(sample))) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱','公园')], 'fake', client=client)
        note = {key:sample[key] for key in ('note_id','content','source_keyword','image_names')}
        await stream.add_note(note)
        await stream.finalize_note(note['note_id'], False, interim=True)
        state = stream.states[(note['note_id'], note['source_keyword'])]
        assert state.status == 'matched' and state.source == 'note'
        assert state.confidence == sample['semantic_confidence']
        assert state.evidence_text == sample['semantic_evidence']
        assert state.target_evidence_text == sample['target_evidence']
        assert state.note['selection_origin'] == 'jev_initial_screen'
        assert state.note['manual_review_required'] is True
        assert state.note['semantic_comment_stop_reason'] == 'body_match'
        assert not state.comments
        await stream.close()
    book=openpyxl.load_workbook(stream.output_path,read_only=True)
    row=dict(zip(next(book['Matched'].values),list(book['Matched'].values)[1]))
    assert row['manual_review_required'] is True
    assert row['jev_original_choice'] in ('keep','review')
    assert dict(book['Summary'].values)['matched'] == 1
    book.close()
    post=browse_posts(tmp_path,stream.output_path)['posts'][0]
    assert post['selection_origin'] == 'jev_initial_screen' and post['manual_review_required']


@pytest.mark.asyncio
async def test_review_with_missing_proof_is_accepted_under_user_policy(tmp_path):
    def respond(request):
        return grounded_response(request,'review',kind='unclear',relation='unclear',confidence=None)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client)
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'公园也许可爱','image_names':'001.jpg'})
        await stream.finalize_note('n',False,interim=True)
        state=stream.states[('n','可爱 公园')]
        assert state.status=='matched' and state.confidence is None
        assert state.note['selection_origin']=='review_policy'
        assert state.note['jev_original_choice']=='review'
        assert state.note['original_semantic_source']=='model_review'
        assert state.note['manual_review_required'] is True
        assert state.note['semantic_comment_stop_reason']=='review_candidate_accepted'
        await stream.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('selection_mode', ['candidate', 'strict'])
@pytest.mark.parametrize('stop_reason', ['comment_budget_exhausted', 'comment_page_budget_exhausted'])
@pytest.mark.parametrize('body,comments', [
    ('这小狗的小辫子太可爱了', ['拍得不错', '好喜欢']),
    ('今天去了公园', ['可爱', '公园']),
    ('今天去了公园', ['路过', '好喜欢']),
])
async def test_comment_limit_excludes_posts_without_a_same_field_keyword_pair(tmp_path, selection_mode, stop_reason, body, comments):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request:pytest.fail('No keyword pair, no model request'))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client,comment_budget=2,selection_mode=selection_mode)
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':body,'image_names':'001.jpg'})
        await stream.finalize_note('n',False,interim=True)
        assert stream.states[('n','可爱 公园')].status is None
        for index, content in enumerate(comments):
            await stream.add_comment({'note_id':'n','comment_id':str(index),'content':content})
            await stream.finalize_note('n',False,interim=True)
            assert stream.states[('n','可爱 公园')].status is None
        await stream.finalize_note('n',False,comment_stop_reason=stop_reason)
        state=stream.states[('n','可爱 公园')]
        assert state.status=='excluded' and state.source=='keyword_miss'
        assert not state.matched and not state.had_candidate
        assert state.quality_reason=='keyword_miss'
        assert state.note['semantic_comment_stop_reason']==stop_reason
        assert state.note['semantic_comments_checked']==2
        assert 'selection_origin' not in state.note and state.evidence_text is None
        assert not state.matched_comment and stream.reviewer.calls==0
        # Finalizing again must not promote an excluded record in close/finally.
        await stream.finalize_note('n',False,comment_stop_reason=stop_reason)
        assert state.status=='excluded'
        await stream.close()
    book=openpyxl.load_workbook(stream.output_path,read_only=True)
    assert book['Matched'].max_row==book['Review'].max_row==1
    assert dict(book['Summary'].values)['excluded']==1
    book.close()
    checkpoint=json.loads(stream.checkpoint_path.read_text(encoding='utf-8').splitlines()[-1])
    assert checkpoint['status']=='excluded' and checkpoint['source']=='keyword_miss'
    assert len(checkpoint['comments'])==2


@pytest.mark.asyncio
@pytest.mark.parametrize('stop_reason', ['comment_budget_exhausted', 'comment_page_budget_exhausted'])
async def test_comment_limit_still_accepts_semantically_uncertain_dual_keyword_candidate(tmp_path, stop_reason):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request:grounded_response(
        request,'review',kind='unclear',relation='unclear',confidence=None))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client,comment_budget=2)
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'今天去了公园','image_names':'001.jpg'})
        await stream.finalize_note('n',False,interim=True)
        await stream.add_comment({'note_id':'n','comment_id':'a','content':'公园也许可爱'})
        await stream.finalize_note('n',False,comment_stop_reason=stop_reason)
        state=stream.states[('n','可爱 公园')]
        assert state.status=='matched' and state.had_candidate and state.matched
        assert state.note['selection_origin']=='review_policy'
        assert state.note['original_semantic_source']==stop_reason
        assert state.note['manual_review_required'] is True
        await stream.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('old_version', ['scene-evidence-v9','scene-evidence-v10','scene-evidence-v11'])
async def test_new_keyword_gate_does_not_reuse_old_acceptance_cache(tmp_path, monkeypatch, old_version):
    import semantic_filter
    current_version=semantic_filter.RULE_VERSION
    assert current_version!=old_version
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request:pytest.fail('No model calls'))) as client:
        monkeypatch.setattr(semantic_filter,'RULE_VERSION',old_version)
        old=StreamingSemanticFilter(tmp_path/'old.xlsx',[('可爱','公园')],'fake',client=client,comment_budget=2)
        await old.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'今天去了公园','image_names':'001.jpg'})
        state=old.states[('n','可爱 公园')]
        state.status,state.source,state.matched='matched','comment_budget_exhausted',True
        state.note['selection_origin']='review_policy'
        old._checkpoint(state)
        old_path=old.checkpoint_path
        old_bytes=old_path.read_bytes()
        await old.close()
        monkeypatch.setattr(semantic_filter,'RULE_VERSION',current_version)
        stream=StreamingSemanticFilter(tmp_path/'new.xlsx',[('可爱','公园')],'fake',client=client,comment_budget=2)
        assert stream.checkpoint_path!=old_path and not stream.states
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'今天去了公园','image_names':'001.jpg'})
        await stream.finalize_note('n',False,comment_stop_reason='comment_budget_exhausted')
        assert stream.states[('n','可爱 公园')].status=='excluded'
        assert old_path.read_bytes()==old_bytes
        await stream.close()


@pytest.mark.asyncio
async def test_review_candidate_target_five_keeps_comment_capture_when_enabled(tmp_path,monkeypatch):
    import config
    from media_platform.weibo import core
    from store.excel_store_base import ExcelStoreBase
    for key,value in {'KEYWORDS':'可爱 公园','START_PAGE':1,'CRAWLER_MAX_NOTES_COUNT':5,
                      'CRAWLER_MAX_SLEEP_SEC':0,'SAVE_DATA_OPTION':'excel','SAVE_DATA_PATH':str(tmp_path),
                      'ENABLE_GET_COMMENTS':True}.items():
        monkeypatch.setattr(config,key,value)
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER',raising=False)
    store=ExcelStoreBase('weibo','search')
    monkeypatch.setenv('MEDIACRAWLER_SEMANTIC_FILTER','1')
    monkeypatch.setattr(ExcelStoreBase,'_instances',{'weibo_search':store})
    monkeypatch.setattr(core.weibo_store.WeibostoreFactory,'create_store',lambda:store)
    monkeypatch.setattr(core,'filter_search_result_card',lambda cards:cards)
    cards=[{'mblog':{'id':str(i),'text':f'公园也许可爱，这是第{i}条记录','created_at':'Thu Oct 08 09:00:00 +0800 2026',
                     'pics':[{'url':'https://wx4.sinaimg.cn/large/test.jpg'}]}} for i in range(1,21)]
    crawler=core.WeiboCrawler()
    crawler.wb_client=SimpleNamespace(get_note_by_keyword=AsyncMock(return_value={'cards':cards}))
    crawler.batch_get_notes_full_text=AsyncMock(side_effect=lambda notes:notes)
    crawler.batch_get_notes_comments=AsyncMock()
    crawler.download_media=AsyncMock()
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: grounded_response(
        request,'review',kind='unclear',relation='unclear',confidence=None))) as client:
        store.semantic_stream=StreamingSemanticFilter(store.filename,[('可爱','公园')],'fake',client=client)
        await crawler.search()
        states=store.semantic_stream.states
        assert len(states)==5 and all(s.status=='matched' for s in states.values())
        assert all(s.note['selection_origin']=='review_policy' for s in states.values())
        assert store.semantic_stream.reviewer.calls==5
        assert crawler.wb_client.get_note_by_keyword.await_count==1
        assert crawler.download_media.await_count==5
        await store.semantic_stream.close()
    store.flush()
    book=openpyxl.load_workbook(store.filename,read_only=True)
    assert book['Contents'].max_row==6
    book.close()
    post=browse_posts(tmp_path,store.filename)['posts'][0]
    assert post['selection_origin'] == 'review_policy' and post['manual_review_required']


@pytest.mark.asyncio
@pytest.mark.parametrize('content,kind,relation',[
    ('带我家可爱的小猫去公园','scene_animal','scene_component'),
    ('公园里吃可爱的蛋糕','food','scene_component'),
    ('带可爱的宝宝去公园','person','scene_component'),
    ('公园真热闹，我买的钥匙扣可爱','unrelated_item','unrelated'),
    ('公园这个英文名字真可爱','place','target_itself'),
    ('公园玩滑梯，昨晚家里看见可爱花坛','landscape','unrelated'),
    ('这个公园一点也不可爱','place','target_itself'),
])
async def test_review_candidates_cannot_bypass_object_or_relation_guards(content,kind,relation):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: grounded_response(
        request,'review',kind=kind,relation=relation,confidence=.25))) as client:
        decision=await JevReviewer('fake',client).decide(('可爱','公园'),content)
    assert decision.choice=='review'


@pytest.mark.asyncio
async def test_primary_export_preserves_candidate_provenance_from_state(tmp_path,monkeypatch):
    import config
    from store.excel_store_base import ExcelStoreBase
    monkeypatch.setattr(config,'SAVE_DATA_PATH',str(tmp_path))
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER',raising=False)
    sample=SAMPLES[1]
    async with httpx.AsyncClient(transport=httpx.MockTransport(replay(sample))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client)
        note={key:sample[key] for key in ('note_id','content','source_keyword','image_names')}
        await stream.add_note(note)
        await stream.finalize_note(note['note_id'],False,interim=True)
        store=ExcelStoreBase('weibo','search')
        store.semantic_stream=stream
        key=(note['note_id'],note['source_keyword'])
        # Buffer predates the decision: provenance exists only in state.note.
        store._semantic_pending_notes[key]=dict(note)
        store._write_retained_semantic_note(note['note_id'])
        store.flush()
        await stream.close()
    book=openpyxl.load_workbook(store.filename,read_only=True)
    row=dict(zip(next(book['Contents'].values),list(book['Contents'].values)[1]))
    assert row['selection_origin']=='jev_initial_screen'
    assert row['manual_review_required'] is True and row['jev_original_choice']=='review'
    book.close()
