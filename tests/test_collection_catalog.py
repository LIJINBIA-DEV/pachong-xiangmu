import hashlib
from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from api.routers import data as router
from collection_tasks import create_task, prepare_run, atomic_json


def save_book(path, semantic=False, retained=('1',), review=('2',)):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.active.title = 'Matched' if semantic else 'Contents'
    for name, ids, status in [(book.active.title, retained, 'matched'), ('Review', review, 'review')]:
        sheet = book[name] if name in book else book.create_sheet(name)
        sheet.append(['note_id', 'title', 'source_keyword', 'semantic_result'])
        for nid in ids:
            sheet.append([nid, '=collected text', '苏州 房价', status])
            sheet.cell(sheet.max_row, 2).data_type = 's'
    comments = book.create_sheet('Comments')
    comments.append(['note_id', 'comment_id', 'content'])
    for nid in (*retained, *review):
        comments.append([nid, 'c' + nid, 'comment ' + nid])
    book.save(path)
    book.close()


@pytest.fixture
def catalog(tmp_path, monkeypatch):
    monkeypatch.setattr(router, 'data_root', lambda: tmp_path)
    app = FastAPI()
    app.include_router(router.router, prefix='/api')
    return tmp_path, TestClient(app)


def attach(root, task, files):
    import json
    path = prepare_run(root, task, {'苏州 房价': 5})
    run = json.loads(path.read_text())
    atomic_json(path, {**run, 'files': files, 'status': 'completed'})


def test_legacy_pair_is_one_task_and_old_round_remains_separate(catalog):
    root, api = catalog
    save_book(root / 'xhs/new.xlsx')
    save_book(root / 'xhs/new_semantic.xlsx', True)
    save_book(root / 'xhs/old_semantic.xlsx', True, (), ())
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob('*.xlsx')}
    tasks = api.get('/api/data/tasks').json()['tasks']
    assert len(tasks) == 2
    latest = next(task for task in tasks if len(task['files']) == 2)
    assert latest['primary_file'] == 'xhs/new.xlsx'
    assert latest['counts'] == {'matched': 1, 'review': 1}
    assert latest['title'] == '苏州 房价'
    old = next(task for task in tasks if len(task['files']) == 1)
    assert old['counts'] == {} and old['primary_file'] == 'xhs/old_semantic.xlsx'
    assert {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before} == before


def test_registered_task_combines_supplements_without_duplicate_posts(catalog):
    root, api = catalog
    task = create_task(root, {}, {'苏州 房价': 5}, anchor='xhs/main.xlsx')
    for name, semantic, retained in [('main.xlsx', False, ('1',)), ('main_semantic.xlsx', True, ('1',)), ('extra.xlsx', False, ('1', '3'))]:
        save_book(root / 'xhs' / name, semantic, retained)
    attach(root, task, ['xhs/main.xlsx', 'xhs/main_semantic.xlsx'])
    attach(root, task, ['xhs/extra.xlsx'])
    result = api.get('/api/data/tasks').json()['tasks']
    assert len(result) == 1
    assert result[0]['id'] == task['id'] and result[0]['rounds'] == 2
    assert result[0]['counts'] == {'matched': 2, 'review': 1}
    browsed = api.get('/api/data/browse', params={'file': 'xhs/main.xlsx', 'status': 'matched'}).json()
    assert browsed['total'] == 2 and browsed['counts'] == result[0]['counts']


def test_missing_or_unreadable_file_does_not_report_false_zero(catalog):
    root, api = catalog
    task = create_task(root, {}, {'苏州 房价': 5}, anchor='xhs/missing.xlsx')
    save_book(root / 'xhs/main_semantic.xlsx', True)
    attach(root, task, ['xhs/missing.xlsx', 'xhs/main_semantic.xlsx'])
    broken = root / 'xhs/broken_semantic.xlsx'
    broken.write_bytes(b'broken')
    tasks = api.get('/api/data/tasks').json()['tasks']
    assert all(task['counts'] is None for task in tasks)
    assert next(item for item in tasks if item['id'] == task['id'])['status'] == 'incomplete'


def test_exports_separate_statuses_and_only_their_comments(catalog):
    root, api = catalog
    path = root / 'xhs/main.xlsx'
    save_book(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    for status, expected in [('matched', '1'), ('review', '2')]:
        response = api.get('/api/data/export', params={'file': 'xhs/main.xlsx', 'status': status})
        assert response.status_code == 200
        book = load_workbook(BytesIO(response.content), data_only=False)
        assert list(book['帖子'].values)[1][0] == expected
        assert len(list(book['帖子'].values)) == 2
        assert list(book['评论'].values)[1][0] == expected
        assert book['帖子']['B2'].data_type == 's'
        book.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    assert api.get('/api/data/export', params={'file': '../outside.xlsx'}).status_code == 403
    assert api.get('/api/data/export', params={'file': 'xhs/main.xlsx', 'status': 'all'}).status_code == 422


def test_manual_conclusion_updates_lists_counts_exports_and_reset_without_rewriting_sources(catalog):
    root, api = catalog
    path = root / 'xhs/main.xlsx'
    save_book(path)
    # The saved body supplies authentic keyword evidence for the review gate.
    book = load_workbook(path)
    for sheet in ('Contents', 'Review'):
        book[sheet].cell(1, 5, 'desc')
        book[sheet].cell(2, 5, '苏州 房价行情')
    book.save(path)
    book.close()
    original = path.read_bytes()
    file = 'xhs/main.xlsx'
    key = '2:苏州 房价'

    def browse(status):
        return api.get('/api/data/browse', params={'file': file, 'status': status}).json()

    def decide(decision):
        result = api.patch('/api/data/review', json={'file': file, 'post_key': key, 'kind': 'post', 'decision': decision, 'partial': True})
        assert result.status_code == 200
        return result.json()['entries'][key]

    assert decide('qualified')['decision_source'] == 'manual'
    assert browse('review')['total'] == 0
    accepted = browse('matched')
    assert accepted['total'] == 2
    confirmed = next(post for post in accepted['posts'] if post['id'] == '2')
    assert confirmed['semantic_status'] == 'review' and confirmed['selection_origin'] == 'user_confirmed'
    assert api.get('/api/data/tasks').json()['tasks'][0]['counts'] == accepted['counts']
    response = api.get('/api/data/export', params={'file': file})
    exported = load_workbook(BytesIO(response.content), read_only=True)
    assert exported['帖子'].max_row == 3 and exported['评论'].max_row == 3
    exported.close()

    decide('unqualified')
    assert browse('review')['total'] == 0
    assert browse('excluded')['posts'][0]['id'] == '2'
    assert browse('matched')['total'] == 1
    decide(None)
    assert browse('review')['total'] == 1
    assert browse('matched')['total'] == 1
    assert path.read_bytes() == original


def test_automatic_image_defaults_and_obsolete_post_labels_do_not_resolve_review(catalog):
    root, api = catalog
    from api.services.manual_review import annotations, write_state
    path = root / 'xhs/main.xlsx'
    save_book(path)
    annotations(root, path)  # Missing images are automatically rejected in image review.
    write_state(root, path, {'version': 1, 'posts': {'2:苏州 房价': {'decision': 'qualified'}}, 'accounts': {}})
    result = api.get('/api/data/browse', params={'file': 'xhs/main.xlsx', 'status': 'review'}).json()
    assert result['total'] == 1 and result['counts'] == {'matched': 1, 'review': 1}


def test_legacy_spatial_export_can_receive_explicit_manual_post_approval(catalog):
    root, api = catalog
    path = root / 'xhs/main.xlsx'
    save_book(path)  # No authentic keyword evidence in the saved content.
    result = api.patch('/api/data/review', json={'file': 'xhs/main.xlsx', 'post_key': '2:苏州 房价', 'kind': 'post', 'decision': 'qualified'})
    assert result.status_code == 200
    posts = api.get('/api/data/browse', params={'file': 'xhs/main.xlsx', 'status': 'matched'}).json()['posts']
    assert any(post['id'] == '2' for post in posts)
