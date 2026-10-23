import json
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl
import pytest

from api.routers import data as data_router
from api.services.data_browser import browse_posts, load_dataset, image_download_info, comment_hits


@pytest.mark.parametrize('names, local_count, expected', [('001.jpg,002.jpg', 0, 'not_downloaded'),
    ('001.jpg,002.jpg', 1, 'partial'), ('001.jpg,002.jpg', 2, 'downloaded'), ('', 0, 'no_images')])
def test_source_images_distinguished_from_downloads(names, local_count, expected):
    images = [{'name': f'{i:03}.jpg', 'local': True} for i in range(1, local_count + 1)]
    info = image_download_info({'image_names': names}, images)
    assert info['image_download_status'] == expected
    assert info['downloaded_image_count'] == local_count
    assert info['image_count'] == (2 if names else 0)


def test_missing_metadata_is_unknown_and_remote_images_are_not_downloaded():
    assert image_download_info({}, []) == {'image_count': None, 'downloaded_image_count': 0, 'image_download_status': 'unknown'}
    info = image_download_info({'image_list': 'https://example.com/1.jpg'}, [{'name': 'image-1', 'local': False}])
    assert info['image_download_status'] == 'not_downloaded' and info['image_count'] == 1


@pytest.mark.parametrize('origin,reason',[('review_policy',None),('user_confirmed',None),('jev_initial_screen','jev_screening_candidate')])
def test_policy_acceptance_does_not_fabricate_verified_comment_hits(origin,reason):
    hits,comments=comment_hits({'semantic_result':'matched','selection_origin':origin,
        'quality_reason':reason,'source_keyword':'可爱 公园','evidence_source':'comment:c'},
        [{'comment_id':'c','content':'公园里的真猫好可爱'}])
    assert hits[0]['kind']=='keyword' and len(comments)==1


@pytest.fixture
def root(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', '')
    monkeypatch.setattr(data_router, 'DATA_DIR', tmp_path)
    return tmp_path


def workbook(root):
    directory = root / 'weibo'
    directory.mkdir(exist_ok=True)
    file = directory / 'weibo_search_test.xlsx'
    book = openpyxl.Workbook()
    notes = book.active
    notes.title = 'Contents'
    notes.append(['note_id', 'title', 'content', 'note_url', 'source_keyword', 'semantic_result', 'evidence_source', 'image_list'])
    notes.append(['1', '猫形花坛', '公园里的猫形花坛真可爱', 'https://m.weibo.cn/detail/1', '可爱 公园', 'matched', 'comment:c1', 'https://example.com/remote.jpg'])
    notes.append(['2', None, '第二个公园', None, '可爱 公园', 'matched', 'content', None])
    review = book.create_sheet('Review')
    review.append(list(next(notes.values)))
    review.append(['3', None, '也许可爱的公园', 'javascript:alert(1)', '可爱 公园', 'review', 'comment:c3', None])
    comments = book.create_sheet('Comments')
    comments.append(['comment_id', 'note_id', 'content'])
    comments.append(['c1', '1', '公园这座花坛真可爱'])
    comments.append(['c2', '2', '公园这个雕塑真可爱'])
    comments.append(['other', '2', '不属于第一帖'])
    reviews = book.create_sheet('ReviewComments')
    reviews.append(['comment_id', 'note_id', 'content'])
    reviews.append(['c3', '3', '公园里的真猫好可爱'])
    book.save(file)
    return file


def test_workbook_joins_only_own_comments_and_distinguishes_hits(root):
    path = workbook(root)
    result = browse_posts(root, path)
    one, two, three = result['posts']
    assert result['counts'] == {'matched': 2, 'review': 1}
    assert one['title'] == '猫形花坛'
    assert [c['id'] for c in one['comments']] == ['c1']
    assert one['hit_comments'][0]['kind'] == 'semantic'
    assert two['hit_comments'][0]['kind'] == 'keyword'
    assert three['hit_comments'][0]['kind'] == 'keyword'
    assert three['url'] == ''
    assert two['title'] == ''  # Do not fabricate titles for Weibo posts.


def test_pagination_search_and_review_filter(root):
    path = workbook(root)
    assert browse_posts(root, path, 1, 1)['posts'][0]['id'] == '2'
    assert browse_posts(root, path, search='不属于第一帖')['posts'][0]['id'] == '2'
    result = browse_posts(root, path, status='review')
    assert result['total'] == 1 and result['file_total'] == 3


def test_local_media_preferred_and_isolated_by_post(root):
    path = workbook(root)
    folder = root / 'wb' / 'media' / '1'
    folder.mkdir(parents=True)
    (folder / '001.jpg').write_bytes(b'image-fixture')
    (folder / 'video.mp4').write_bytes(b'video-fixture')
    posts = browse_posts(root, path)['posts']
    assert posts[0]['images'] == [{'url': '/api/data/media/wb/media/1/001.jpg', 'name': '001.jpg', 'local': True}]
    assert posts[1]['images'] == []


def test_cached_workbook_refreshes_after_new_results(root):
    path = workbook(root)
    assert len(load_dataset(path)[0]) == 3
    book = openpyxl.load_workbook(path)
    book['Contents'].append(['4', '新增', '内容', None, '', 'matched'])
    book.save(path)
    assert len(load_dataset(path)[0]) == 4


def test_json_pairs_comments_and_jsonl_latest_state(root):
    directory = root / 'xhs'
    directory.mkdir()
    notes = directory / 'search_contents_test.json'
    notes.write_text(json.dumps([{'note_id': 'a', 'title': '标题', 'desc': '内容', 'source_keyword': '可爱 公园', 'image_list': 'https://example.com/1.jpg,https://example.com/2.jpg'}]), encoding='utf-8')
    (directory / 'search_comments_test.json').write_text(json.dumps([{'comment_id': 'c', 'note_id': 'a', 'content': '可爱的公园'}]), encoding='utf-8')
    result = browse_posts(root, notes)['posts'][0]
    assert len(result['images']) == 2 and result['hit_comments'][0]['kind'] == 'keyword'
    checkpoint = directory / 'saved.jsonl'
    rows = [{'note': {'note_id': 'b', 'content': '正文'}, 'status': status, 'comments': []} for status in ['pending', 'matched']]
    checkpoint.write_text('\n'.join(json.dumps(row) for row in rows), encoding='utf-8')
    assert browse_posts(root, checkpoint)['posts'][0]['status'] == 'matched'


def client():
    app = FastAPI()
    app.include_router(data_router.router, prefix='/api')
    return TestClient(app)


def test_api_file_list_normalized_and_query_validation(root):
    workbook(root)
    api = client()
    file = api.get('/api/data/files').json()['files'][0]
    assert '\\' not in file['path']
    result = api.get('/api/data/browse', params={'file': file['path']})
    assert result.status_code == 200 and result.json()['total'] == 3
    assert api.get('/api/data/browse', params={'file': file['path'], 'offset': -1}).status_code == 422


def test_api_prevents_traversal_and_non_image_access(root):
    outside = root.parent / 'outside.json'
    outside.write_text('[]')
    api = client()
    assert api.get('/api/data/browse', params={'file': '../outside.json'}).status_code == 403
    assert api.get('/api/data/browse', params={'file': str(outside)}).status_code == 403
    folder = root / 'wb' / 'media' / '1'
    folder.mkdir(parents=True)
    (folder / '001.png').write_bytes(b'image-fixture')
    (folder / 'private.txt').write_text('not an image')
    assert api.get('/api/data/media/wb/media/1/001.png').status_code == 200
    assert api.get('/api/data/media/wb/media/1/private.txt').status_code == 403


def test_configured_data_directory_matches_crawler(root, monkeypatch):
    import config
    configured = root / 'custom'
    configured.mkdir()
    path = workbook(configured)
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(configured))
    api = client()
    assert api.get('/api/data/files').json()['files'][0]['path'] == path.relative_to(configured).as_posix()
    assert api.get('/api/data/browse', params={'file': 'weibo/weibo_search_test.xlsx'}).status_code == 200
