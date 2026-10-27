from datetime import datetime
import json

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from api.routers import data as router_module


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(router_module, 'data_root', lambda: tmp_path)
    app = FastAPI()
    app.include_router(router_module.router, prefix='/api')
    return TestClient(app)


def test_preview_selects_nonempty_sheet_and_can_switch_to_empty_sheet(tmp_path, client):
    book = Workbook()
    book.active.title = 'Matched'
    book.active.append(['note_id', 'title'])
    book.create_sheet('Pending').append(['note_id', 'title'])
    summary = book.create_sheet('Summary')
    summary.append(['matched', 0])
    summary.append(['excluded', 35])
    book.save(tmp_path/'semantic.xlsx')

    response = client.get('/api/data/files/semantic.xlsx')
    assert response.status_code == 200
    result = response.json()
    assert result['sheet'] == 'Summary'
    assert result['data'] == [{'metric': 'matched', 'value': 0}, {'metric': 'excluded', 'value': 35}]
    assert result['sheets'] == ['Matched', 'Pending', 'Summary']
    empty = client.get('/api/data/files/semantic.xlsx?sheet=Matched').json()
    assert empty['total'] == 0 and empty['columns'] == ['note_id', 'title']
    assert client.get('/api/data/files/semantic.xlsx?sheet=Missing').status_code == 400


def test_preview_handles_blank_numeric_cells_dates_and_limits(tmp_path, client):
    book = Workbook()
    sheet = book.active
    sheet.append(['note_id', 'likes', 'saved_at'])
    sheet.append(['1', 10, datetime(2026, 10, 26)])
    sheet.append(['2', None, None])
    book.save(tmp_path/'posts.xlsx')
    response = client.get('/api/data/files/posts.xlsx')
    assert response.status_code == 200
    rows = response.json()['data']
    assert rows[1]['likes'] is None
    assert rows[0]['saved_at'] == '2026-10-26T00:00:00'
    limited = client.get('/api/data/files/posts.xlsx?limit=1').json()
    assert len(limited['data']) == 1 and limited['total'] == 2
    assert client.get('/api/data/files/posts.xlsx?limit=0').status_code == 422


def test_json_object_is_returned_as_table_rows(tmp_path, client):
    (tmp_path/'record.json').write_text(json.dumps({'note_id': '1', 'title': 'test'}))
    assert client.get('/api/data/files/record.json').json()['data'] == [{'note_id': '1', 'title': 'test'}]
