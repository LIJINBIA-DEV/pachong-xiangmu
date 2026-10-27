import hashlib

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from api.routers import data as router


@pytest.mark.parametrize(('semantic', 'retained', 'review'), [(False, 5, 2), (True, 5, 2), (True, 0, 0)])
def test_file_list_reports_post_counts_without_counting_comments_or_headers(tmp_path, monkeypatch, semantic, retained, review):
    monkeypatch.setattr(router, 'data_root', lambda: tmp_path)
    path = tmp_path / ('xhs_search_semantic.xlsx' if semantic else 'xhs_search.xlsx')
    book = Workbook()
    book.active.title = 'Matched' if semantic else 'Contents'
    book.active.append(['note_id', 'title'])
    for index in range(retained):
        book.active.append([str(index), 'retained'])
    # Styled blank rows must not inflate the post count.
    book.active.cell(20, 1).number_format = '@'
    pending = book.create_sheet('Review')
    pending.append(['note_id', 'title'])
    for index in range(review):
        pending.append([str(index), 'review'])
    comments = book.create_sheet('MatchedComments' if semantic else 'Comments')
    comments.append(['comment_id', 'note_id'])
    for index in range(40):
        comments.append([str(index), '1'])
    book.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    app = FastAPI()
    app.include_router(router.router, prefix='/api')
    response = TestClient(app).get('/api/data/files')
    assert response.status_code == 200
    info = response.json()['files'][0]
    assert info['file_kind'] == ('semantic' if semantic else 'collection')
    assert info['retained_count'] == retained
    assert info['review_count'] == review
    assert info['name'] == path.name
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest


def test_unreadable_workbook_has_unknown_counts_instead_of_zero(tmp_path, monkeypatch):
    monkeypatch.setattr(router, 'data_root', lambda: tmp_path)
    path = tmp_path/'broken_semantic.xlsx'
    path.write_bytes(b'not a valid workbook')
    info = router.get_file_info(path)
    assert info['file_kind'] == 'semantic'
    assert info['retained_count'] is None and info['review_count'] is None
