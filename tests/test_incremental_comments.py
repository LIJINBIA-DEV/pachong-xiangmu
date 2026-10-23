"""Exercise text-first semantic review, bounded pagination, and evidence exports."""
import asyncio
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import openpyxl
import pytest

from semantic_filter import Decision, StreamingSemanticFilter
from tests.semantic_responses import grounded_response


@pytest.mark.asyncio
async def test_body_review_is_provisional_and_new_comment_can_rescue_it(tmp_path):
    sent = []
    def respond(request):
        payload = json.loads(request.content)
        sent.append(payload)
        if '一级评论 hit' in payload['state']['text']:
            return grounded_response(request, source='comment:hit')
        return grounded_response(request, 'review', relation='unclear')
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client, selection_mode='strict')
        await stream.add_note({'note_id': 'n', 'source_keyword': '可爱 公园', 'content': '这个公园也许可爱', 'image_names': '001.jpg'})
        await stream.finalize_note('n', False, interim=True)
        state = stream.states[('n', '可爱 公园')]
        assert state.status is None and state.uncertain
        await stream.add_comment({'note_id': 'n', 'comment_id': 'old', 'content': '第一批无关评论'})
        await stream.finalize_note('n', False, interim=True)
        assert state.status is None and state.reviewed_comment_count == 1
        await stream.add_comment({'note_id': 'n', 'comment_id': 'hit', 'content': '公园这个花坛真可爱'})
        await stream.finalize_note('n', False, interim=True)
        assert state.status == 'matched' and state.source == 'comment'
        assert state.matched_comment['comment_id'] == 'hit'
        assert state.evidence_source == 'comment:hit'
        assert '第一批无关评论' not in sent[-1]['state']['text']
        calls = stream.reviewer.calls
        await stream.finalize_note('n', False, comment_stop_reason='evidence_found')
        assert stream.reviewer.calls == calls  # Do not judge again in finally/close.
        await stream.close()
    book = openpyxl.load_workbook(stream.output_path, read_only=True)
    assert book['Review'].max_row == 1
    assert list(book['MatchedComments'].values)[1][1] == 'hit'
    book.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('body_choice', ['reject', 'review'])
async def test_budget_keeps_only_uncertain_candidates_without_rejudging(tmp_path, body_choice):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: grounded_response(request, body_choice, relation='unclear'))) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client, comment_budget=2, selection_mode='strict')
        await stream.add_note({'note_id': 'n', 'source_keyword': '可爱 公园', 'content': '在公园看到可爱的小狗', 'image_names': '001.jpg'})
        await stream.finalize_note('n', False, interim=True)
        await stream.add_comment({'note_id': 'n', 'comment_id': 'a', 'content': '拍得不错'})
        await stream.finalize_note('n', False, interim=True)
        calls = stream.reviewer.calls
        await stream.finalize_note('n', False, comment_stop_reason='comment_budget_exhausted')
        state = stream.states[('n', '可爱 公园')]
        assert state.status == ('excluded' if body_choice == 'reject' else 'review')
        assert state.source == ('semantic_reject' if body_choice == 'reject' else 'comment_budget_exhausted')
        assert state.comments_available is False
        assert stream.reviewer.calls == calls
        await stream.close()


@pytest.mark.asyncio
async def test_full_comment_exhaustion_and_zero_budget_both_exclude_keyword_miss(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client, selection_mode='strict')
        for nid in ('complete', 'budget'):
            await stream.add_note({'note_id': nid, 'source_keyword': '可爱 公园', 'content': '今天去公园', 'image_names': '001.jpg'})
            await stream.finalize_note(nid, False, interim=True)
            await stream.finalize_note(nid, nid == 'complete', comment_stop_reason='comments_exhausted' if nid == 'complete' else 'comment_budget_exhausted')
        assert stream.states[('complete', '可爱 公园')].source == 'keyword_miss'
        assert stream.states[('budget', '可爱 公园')].status == 'excluded'
        assert stream.states[('budget', '可爱 公园')].source == 'keyword_miss'
        assert stream.reviewer.calls == 0
        await stream.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('limit, cursor, expected_complete', [(2, 0, False), (3, 0, True), (3, 8, False), (0, 0, False)])
async def test_client_reports_budget_even_when_server_final_page_is_truncated(limit, cursor, expected_complete):
    from media_platform.weibo.client import WeiboClient
    client = object.__new__(WeiboClient)
    client.get_note_comments = AsyncMock(return_value={'max_id': cursor, 'data': [{'id': str(i)} for i in range(3)]})
    completed = AsyncMock()
    page = AsyncMock(return_value=False)
    comments = await client.get_note_all_comments('n', 0, max_count=limit, include_sub_comments=False, page_callback=page, completion_callback=completed)
    assert len(comments) == limit
    assert client.get_note_comments.await_count == (0 if limit == 0 else 1)
    completed.assert_awaited_once_with(expected_complete, 'comments_exhausted' if expected_complete else 'comment_budget_exhausted')


@pytest.mark.asyncio
async def test_client_early_stop_has_no_next_request_or_unneeded_sleep(monkeypatch):
    from media_platform.weibo import client as module
    client = object.__new__(module.WeiboClient)
    client.get_note_comments = AsyncMock(return_value={'max_id': 9, 'data': [{'id': 'hit'}]})
    sleep = AsyncMock()
    monkeypatch.setattr(module.asyncio, 'sleep', sleep)
    complete = AsyncMock()
    await client.get_note_all_comments('n', 2, max_count=10, include_sub_comments=False, page_callback=AsyncMock(return_value=True), completion_callback=complete)
    assert client.get_note_comments.await_count == 1
    sleep.assert_not_awaited()
    complete.assert_awaited_once_with(False, 'evidence_found')


@pytest.mark.asyncio
async def test_empty_pages_have_a_budget_and_repeated_comments_do_not_use_count_budget():
    from media_platform.weibo.client import WeiboClient
    client = object.__new__(WeiboClient)
    client.get_note_comments = AsyncMock(side_effect=[{'max_id': 1, 'data': []}, {'max_id': 2, 'data': []}])
    complete = AsyncMock()
    await client.get_note_all_comments('n', 0, max_count=10, include_sub_comments=False, max_pages=2, completion_callback=complete)
    complete.assert_awaited_once_with(False, 'comment_page_budget_exhausted')
    client.get_note_comments = AsyncMock(side_effect=[{'max_id': 1, 'data': [{'id': 'a'}]}, {'max_id': 0, 'data': [{'id': 'a'}, {'id': 'b'}]}])
    comments = await client.get_note_all_comments('n', 0, max_count=2, include_sub_comments=False)
    assert [item['id'] for item in comments] == ['a', 'b']


@pytest.mark.asyncio
async def test_comments_budget_changes_checkpoint_identity(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        first = StreamingSemanticFilter(tmp_path / 'one.xlsx', [('可爱', '公园')], 'fake', client=client, comment_budget=10)
        await first.add_note({'note_id': 'n', 'source_keyword': '可爱 公园', 'content': '公园', 'image_names': '001.jpg'})
        await first.finalize_note('n', False, comment_stop_reason='comment_budget_exhausted')
        await first.close()
        second = StreamingSemanticFilter(tmp_path / 'two.xlsx', [('可爱', '公园')], 'fake', client=client, comment_budget=50)
        assert second.checkpoint_path != first.checkpoint_path and not second.states
        await second.close()


@pytest.mark.asyncio
async def test_service_failure_preserves_complete_fetched_comments_for_resume(tmp_path, monkeypatch):
    import semantic_filter
    from store.excel_store_base import ExcelStoreBase
    monkeypatch.setattr(semantic_filter.asyncio, 'sleep', AsyncMock())
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: httpx.Response(503))) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client)
        await stream.add_note({'note_id': 'n', 'source_keyword': '可爱 公园', 'content': '公园', 'image_names': '001.jpg'})
        await stream.finalize_note('n', False, interim=True)
        await stream.add_comment({'note_id': 'n', 'comment_id': 'a', 'content': '这个公园好可爱'})
        instance = SimpleNamespace(semantic_stream=stream, _write_retained_semantic_note=lambda nid: None)
        monkeypatch.setattr(ExcelStoreBase, '_instances', {'weibo_search': instance})
        assert await ExcelStoreBase.review_semantic_note('weibo', 'n') == 'semantic_service_pending'
        await stream.finalize_note('n', True, comment_stop_reason='semantic_service_pending')
        assert stream.states[('n', '可爱 公园')].status == 'pending'
        assert stream.states[('n', '可爱 公园')].comments_available is True
        await stream.close()


def test_web_comment_budget_reaches_crawler_command():
    from api.schemas.crawler import CrawlerStartRequest
    from api.services.crawler_manager import CrawlerManager
    request = CrawlerStartRequest(platform='wb', max_comments_count=50)
    cmd = CrawlerManager()._build_command(request)
    assert cmd[cmd.index('--max_comments_count_singlenotes') + 1] == '50'


@pytest.mark.asyncio
@pytest.mark.parametrize('budget, expected_status, checked', [(10, 'review', 10), (50, 'matched', 40)])
async def test_real_622_comment_sample_replays_saved_evidence_without_fetching_to_end(tmp_path, budget, expected_status, checked):
    # Reuse the recorded verified decision when its original evidence arrives.
    # This checks pagination/evidence preservation, not current Jev accuracy.
    fixture = json.loads((Path(__file__).parent / 'fixtures/weibo_incremental_20261008.json').read_text(encoding='utf-8'))
    note, evidence = fixture['note'], fixture['verified_decision']
    assert fixture['original_collected_count'] == 622 and fixture['matched_position'] == 30
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: pytest.fail('Replay must not use a remote model'))) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client, comment_budget=budget, selection_mode='strict')
        async def replay(_pair, _text, _ids, records):
            if any(item.source == evidence['evidence_source'] and evidence['evidence_text'] in item.text for item in records):
                return Decision('keep', **evidence)
            return Decision('review')
        stream.reviewer.decide = replay
        await stream.add_note(note)
        await stream.finalize_note(note['note_id'], False, interim=True)
        assert stream.states[(note['note_id'], note['source_keyword'])].status is None
        for offset in range(0, min(budget, len(fixture['comments'])), 20):
            for comment in fixture['comments'][offset:min(offset + 20, budget)]:
                await stream.add_comment(comment)
            await stream.finalize_note(note['note_id'], False, interim=True)
            if stream.states[(note['note_id'], note['source_keyword'])].status == 'matched':
                break
        await stream.finalize_note(note['note_id'], False, comment_stop_reason='comment_budget_exhausted')
        state = stream.states[(note['note_id'], note['source_keyword'])]
        assert state.status == expected_status and len(state.comments) == checked
        if expected_status == 'matched':
            assert state.matched_comment['comment_id'] == '5351318216771896'
            assert state.evidence_source == evidence['evidence_source']
        await stream.close()


@pytest.mark.asyncio
async def test_live_search_body_first_comment_match_budget_and_export(tmp_path, monkeypatch):
    from media_platform.weibo import core
    from media_platform.weibo.client import WeiboClient
    from store.excel_store_base import ExcelStoreBase
    import config
    for name, value in {'KEYWORDS': '可爱 公园', 'START_PAGE': 1, 'CRAWLER_MAX_NOTES_COUNT': 10,
                        'CRAWLER_MAX_SLEEP_SEC': 0, 'CRAWLER_MAX_COMMENTS_COUNT_SINGLENOTES': 3,
                        'ENABLE_GET_COMMENTS': True, 'SAVE_DATA_OPTION': 'excel', 'SAVE_DATA_PATH': str(tmp_path)}.items():
        monkeypatch.setattr(config, name, value)
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER', raising=False)
    store = ExcelStoreBase('weibo', 'search')
    monkeypatch.setenv('MEDIACRAWLER_SEMANTIC_FILTER', '1')
    monkeypatch.setattr(ExcelStoreBase, '_instances', {'weibo_search': store})
    monkeypatch.setattr(core.weibo_store.WeibostoreFactory, 'create_store', lambda: store)
    monkeypatch.setattr(core, 'filter_search_result_card', lambda cards: cards)
    texts = {'empty': '公园好可爱', 'body': '公园里的花坛真可爱', 'comment': '今天去公园', 'budget': '今天去公园', 'exhausted': '今天去公园'}
    cards = [{'mblog': {'id': nid, 'text': content, 'created_at': 'Thu Oct 08 09:00:00 +0800 2026',
                         'pics': [] if nid == 'empty' else [{'url': 'https://wx4.sinaimg.cn/large/test.jpg'}]}} for nid, content in texts.items()]
    client = object.__new__(WeiboClient)
    client.get_note_by_keyword = AsyncMock(side_effect=[{'cards': cards}, {'cards': []}])
    calls = []
    def comment(cid, text):
        return {'id': cid, 'text': text, 'created_at': 'Thu Oct 08 09:00:00 +0800 2026'}
    async def get_page(nid, cursor, cursor_type):
        calls.append((nid, cursor))
        if nid == 'comment':
            if cursor == -1:
                return {'max_id': 1, 'data': [comment('unrelated', '第一批无关评论')]}
            return {'max_id': 0, 'data': [comment('hit', '公园里的花坛真可爱')]}
        if nid == 'body':
            return {'max_id': 0, 'data': [comment('body-extra', '确实很可爱')]}
        if nid == 'budget':
            return {'max_id': 8, 'data': [comment(str(i), '路过') for i in range(20)]}
        if nid == 'exhausted':
            return {'max_id': 0, 'data': []}
        raise AssertionError('No comments should be fetched for pictureless or excluded posts')
    client.get_note_comments = get_page
    crawler = core.WeiboCrawler()
    crawler.wb_client = client
    crawler.batch_get_notes_full_text = AsyncMock(side_effect=lambda notes: notes)
    crawler.download_media = AsyncMock()
    def respond(request):
        payload = json.loads(request.content)
        return grounded_response(request, source='comment:hit' if '一级评论 hit' in payload['state']['text'] else None)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as model_client:
        store.semantic_stream = StreamingSemanticFilter(store.filename, [('可爱', '公园')], 'fake', client=model_client, comment_budget=3)
        await crawler.search()
        states = store.semantic_stream.states
        assert states[('empty', '可爱 公园')].source == 'no_images'
        assert states[('body', '可爱 公园')].status == 'matched'
        assert states[('comment', '可爱 公园')].evidence_source == 'comment:hit'
        assert states[('budget', '可爱 公园')].status == 'excluded'
        assert states[('budget', '可爱 公园')].source == 'keyword_miss'
        assert states[('budget', '可爱 公园')].note['semantic_comment_stop_reason'] == 'comment_budget_exhausted'
        assert states[('exhausted', '可爱 公园')].status == 'excluded'
        assert len(states[('budget', '可爱 公园')].comments) == 3
        assert store.semantic_stream.reviewer.calls == 4
        await store.semantic_stream.close()
    assert calls == [('body', -1), ('comment', -1), ('comment', 1), ('budget', -1), ('exhausted', -1)]
    assert [call.args[0]['id'] for call in crawler.download_media.await_args_list] == ['body', 'comment']
    store.flush()
    book = openpyxl.load_workbook(store.filename, read_only=True)
    def rows(name):
        iterator = iter(book[name].values)
        headers = next(iterator)
        return [dict(zip(headers, row)) for row in iterator]
    contents = rows('Contents')
    assert [item['note_id'] for item in contents] == ['body', 'comment']
    assert [item['comment_id'] for item in rows('Comments')] == ['body-extra', 'unrelated', 'hit']
    assert contents[0]['comments_collected_count'] == 1
    assert contents[0]['comments_collection_complete'] is True
    book.close()
