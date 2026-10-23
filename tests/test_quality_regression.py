"""Real input regression plus controlled service failures. No model accuracy claim."""

import json
from pathlib import Path
from types import SimpleNamespace

import httpx
import openpyxl
import pytest

from semantic_filter import Decision, JevReviewer, LiveNote, StreamingSemanticFilter, _evidence_chunks, filter_checkpoint
from semantic_quality import Evidence, evidence_records, keep_problem, duplicate_kind, scene_animal_ambiguity
from tests.semantic_responses import grounded_response

DATA = json.loads((Path(__file__).parent / "fixtures/weibo_quality_20260929.json").read_text(encoding="utf-8"))
CASES = {row["id"]: row for row in DATA["cases"]}


def records_for(row):
    fields = [("content", row["note"]["content"])]
    fields.extend((f"comment:{c['comment_id']}", c["content"]) for c in row["comments"])
    return evidence_records(fields)


def selected_decision(row, *, confidence=None, evidence_text=None):
    old = row["previous"]
    quote = evidence_text if evidence_text is not None else old["evidence_text"] or ""
    for prefix in ("笔记上下文（以下标题、正文、一级评论属于同一条微博）：", "content: "):
        quote = quote.removeprefix(prefix)
    records = records_for(row)
    selected = next((record for record in records if quote.strip() in record.text and quote.strip()), None)
    location = next((record for record in records if "公园" in record.text), None)
    return Decision(
        "keep", old["confidence"] if confidence is None else confidence,
        object_type=old["object_type"], target_relation=old["target_relation"],
        evidence_id=selected.key if selected else None,
        target_evidence_id=location.key if location else None,
    ), records


@pytest.mark.parametrize("nid", ["5348427791861259", "5347896161731979", "5348198942249797"])
def test_recorded_weak_keeps_cannot_pass(nid):
    decision, records = selected_decision(CASES[nid])
    assert keep_problem(("可爱", "公园"), decision, records, .8) is not None


def test_name_evaluation_stays_invalid_even_with_high_confidence():
    decision, records = selected_decision(CASES["5347896161731979"], confidence=.99)
    assert keep_problem(("可爱", "公园"), decision, records, .8) == "name_only_evidence"


def test_unrelated_comment_evidence_stays_invalid_even_with_high_confidence():
    decision, records = selected_decision(CASES["5348198942249797"], confidence=.99)
    decision.target_relation = "scene_component"
    assert keep_problem(("可爱", "公园"), decision, records, .8) == "evaluation_without_keyword"


@pytest.mark.parametrize("nid", ["5348459936746114", "5348262121046076", "5348096089526688"])
def test_recorded_animal_rejections_require_reassessment(nid):
    row = CASES[nid]
    assert row["previous"]["status"] == "excluded"
    assert row["previous"]["object_type"] == "animal"
    assert scene_animal_ambiguity(row["note"]["content"])


@pytest.mark.parametrize("nid", ["5347834165461383", "5348451778826261", "5345287497319914"])
def test_real_animals_and_unrelated_items_do_not_trigger_scene_rescue(nid):
    assert not scene_animal_ambiguity(CASES[nid]["note"]["content"])


def test_real_exact_and_near_copies_but_not_independent_park_experiences():
    original = CASES["5348469663073464"]
    previous = SimpleNamespace(note=original["note"], source="note", evidence_text=original["previous"]["evidence_text"])
    assert duplicate_kind(CASES["5348452051717984"]["note"], "note", None, [previous]) == ("exact_text", "5348469663073464")
    long_news = SimpleNamespace(note=CASES["5348068222569297"]["note"], source="note", evidence_text=None)
    assert duplicate_kind(CASES["5348220331886571"]["note"], "note", None, [long_news]) == ("near_text", "5348068222569297")
    personal = SimpleNamespace(note=CASES["5348212611484438"]["note"], source="note", evidence_text=None)
    assert duplicate_kind(CASES["5348103592871413"]["note"], "note", None, [long_news, previous, personal]) is None


def test_identical_short_captions_and_distinct_comments_are_not_discarded():
    old = SimpleNamespace(note={"note_id": "1", "content": "这个公园好可爱"}, source="note", evidence_text="这个公园好可爱")
    assert duplicate_kind({"note_id": "2", "content": "这个公园好可爱"}, "note", "", [old]) is None
    old.note["content"] = CASES["5348469663073464"]["note"]["content"]
    old.source = "comment"
    old.evidence_text = "这个公园的花坛好可爱"
    assert duplicate_kind(old.note, "comment", "这个公园的门头好可爱", [old]) is None


def test_all_evidence_is_original_source_slice_and_long_text_is_covered():
    for row in CASES.values():
        state = LiveNote(row["note"], ("可爱", "公园"), comments=row["comments"])
        sources = {"content": row["note"]["content"], **{f"comment:{c['comment_id']}": c["content"] for c in row["comments"]}}
        covered = {source: set() for source in sources}
        for text, _, records in _evidence_chunks(state):
            assert len(text) <= 3000
            for record in records:
                assert record.text == sources[record.source][record.start:record.end]
                covered[record.source].update(range(record.start, record.end))
        for source, content in sources.items():
            assert all(i in covered[source] for i, char in enumerate(content) if not char.isspace())


@pytest.mark.asyncio
@pytest.mark.parametrize("rescue_succeeds", [True, False])
async def test_explicit_animal_rejection_is_not_rechecked_into_acceptance(rescue_succeeds):
    text = CASES["5348459936746114"]["note"]["content"]
    sent = []
    def respond(request):
        payload = json.loads(request.content)
        sent.append(payload)
        if len(sent) == 1 or not rescue_succeeds:
            return grounded_response(request, "reject", kind="animal", relation="unrelated")
        return grounded_response(request, kind="landscape", relation="scene_component")
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        reviewer = JevReviewer("fake", client)
        decision = await reviewer.decide(("可爱", "公园"), text)
    assert decision.choice == "reject" and decision.reason == "semantic_reject"
    assert len(sent) == 1 and not decision.reviewed_twice


@pytest.mark.asyncio
async def test_source_is_taken_from_selected_comment_not_model_source_claim(tmp_path):
    def respond(request):
        result = grounded_response(request, source="comment:c1", kind="landscape", relation="scene_component")
        data = result.json()
        data["answers"]["evidence"] = {"choice": "post"}  # obsolete contradictory claim
        return httpx.Response(200, json=data)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        await stream.add_note({"note_id": "1", "content": "今天去了公园", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await stream.add_comment({"note_id": "1", "comment_id": "c1", "content": "公园的兔子装置真可爱"})
        await stream.finalize_note("1")
        await stream.close()
    state = stream.states[("1", "可爱 公园")]
    assert state.status == "matched"
    assert state.source == "comment" and state.evidence_source == "comment:c1"
    assert state.matched_comment["comment_id"] == "c1"
    result = openpyxl.load_workbook(stream.output_path)
    assert list(result["MatchedComments"].values)[1][1] == "c1"


@pytest.mark.asyncio
async def test_high_confidence_but_contradicted_evidence_is_rejected():
    text = CASES["5348427791861259"]["note"]["content"]
    def respond(request):
        payload = json.loads(request.content)
        if "support" in payload["questions"]:
            return httpx.Response(200, json={"answers": {"support": {"choice": "contradicted", "confidence": .99}}})
        return grounded_response(request, kind="facility", relation="scene_component", confidence=.99)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        decision = await JevReviewer("fake", client).decide(("可爱", "公园"), text)
    assert decision.choice == "reject" and decision.reason == "evidence_contradicted"


@pytest.mark.asyncio
async def test_exact_duplicates_do_not_count_and_survive_resume(tmp_path):
    ids = ["5348469663073464", "5348452051717984"]
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        for nid in ids:
            await stream.add_note({**CASES[nid]["note"], "image_names": "001.jpg"})
            await stream.finalize_note(nid)
        assert stream.states[(ids[0], "可爱 公园")].status == "matched"
        assert stream.states[(ids[1], "可爱 公园")].status == "duplicate"
        assert stream.states[(ids[1], "可爱 公园")].duplicate_of == ids[0]
        await stream.close()
        resumed = StreamingSemanticFilter(tmp_path / "again.xlsx", [("可爱", "公园")], "fake", client=client)
        assert resumed.states[(ids[1], "可爱 公园")].status == "duplicate"
        await resumed.close()
    result = openpyxl.load_workbook(stream.output_path)
    summary = dict(result["Summary"].values)
    assert summary["matched"] == 1 and summary["duplicate"] == 1
    assert result["Duplicates"].max_row == 2


@pytest.mark.asyncio
@pytest.mark.parametrize("row", [r for r in CASES.values() if r["case"] in {"positive", "negative"}], ids=lambda r: r["id"])
async def test_real_positive_and_negative_inputs_with_controlled_model_decisions(tmp_path, row):
    """Validates guards on real text; the model decision is explicitly controlled."""
    positive = row["case"] == "positive"
    def respond(request):
        return grounded_response(request, "keep" if positive else "reject", kind="landscape" if positive else "animal", relation="scene_component" if positive else "unrelated")
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        await stream.add_note({**row["note"], "image_names": "001.jpg"})
        for comment in row["comments"]:
            await stream.add_comment({**comment, "note_id": row["id"]})
        await stream.finalize_note(row["id"])
        assert stream.states[(row["id"], "可爱 公园")].status == ("matched" if positive else "excluded")
        await stream.close()


@pytest.mark.parametrize("confidence", [None, .2, True, float("nan"), 1.1])
def test_invalid_or_low_confidence_cannot_pass(confidence):
    decision, records = selected_decision(CASES["5347896161731979"])
    decision.confidence = confidence
    assert keep_problem(("可爱", "公园"), decision, records, .8) == "low_or_missing_confidence"


@pytest.mark.asyncio
async def test_verification_service_failure_stays_pending(tmp_path):
    def respond(request):
        if "support" in json.loads(request.content)["questions"]:
            return httpx.Response(200, json={"answers": {"support": "malformed"}})
        return grounded_response(request)
    from semantic_filter import SemanticFilterError
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        await stream.add_note({"note_id": "1", "source_keyword": "可爱 公园", "content": "这个公园真可爱", "image_names": "001.jpg"})
        with pytest.raises(SemanticFilterError):
            await stream.finalize_note("1")
        assert stream.states[("1", "可爱 公园")].status == "pending"
        await stream.close()


@pytest.mark.asyncio
async def test_offline_output_isolated_and_source_preserved(tmp_path):
    row = CASES["5348431807644308"]
    source = tmp_path / "old.jsonl"
    original = json.dumps({"note": row["note"], "comments": row["comments"], "comments_available": True})
    source.write_text(original, encoding="utf-8")
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        output, counts = await filter_checkpoint(source, [("可爱", "公园")], "fake", client=client, output_dir=tmp_path / "new")
    assert output.parent == tmp_path / "new"
    assert source.read_text(encoding="utf-8") == original
    assert counts["matched"] == 1  # Missing historical metadata is now a manual-review candidate.
    assert len(list((tmp_path / "new").glob("semantic_checkpoint_*.jsonl"))) == 1
