"""Verify real downloaded packages through the review API."""
from io import BytesIO
import json
import time
import pytest
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from api.routers import data as router
from api.services.manual_review import update_annotation, read_state


def ready(client, job):
    for _ in range(150):
        job = client.get('/api/data/review/export/jobs/' + job['id']).json()
        if job['status'] not in ('queued', 'preparing'):
            return job
        time.sleep(.02)
    raise AssertionError(job)


def test_user_downloads_excel_and_approved_original_images(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    path = tmp_path / 'weibo/posts.json'
    path.parent.mkdir()
    path.write_text(json.dumps([
        {'note_id': 'one', 'content': '=literal', 'image_names': 'good.jpg,bad.jpg'},
        {'note_id': 'two', 'image_names': 'pending.jpg'},
    ]), encoding='utf-8')
    media = tmp_path / 'wb/media/one'
    media.mkdir(parents=True)
    for name in ('good.jpg', 'bad.jpg'):
        Image.new('RGB', (20, 30), 'green').save(media / name)
    update_annotation(tmp_path, path, 'one:', 'image', 'qualified', 'good.jpg')
    update_annotation(tmp_path, path, 'one:', 'image', 'unqualified', 'bad.jpg')
    update_annotation(tmp_path, path, 'two:', 'image', None, 'pending.jpg')
    update_annotation(tmp_path, path, 'two:', 'post', None)
    before = path.read_bytes(), read_state(tmp_path, path)
    app = FastAPI()
    app.include_router(router.router, prefix='/api')
    with TestClient(app) as client:
        response = client.post('/api/data/review/export/jobs', json={'file': 'weibo/posts.json'})
        assert response.status_code == 200, response.text
        job = response.json()
        for _ in range(100):
            job = client.get('/api/data/review/export/jobs/' + job['id']).json()
            if job['status'] not in ('queued', 'preparing'):
                break
            time.sleep(.02)
        assert job['status'] == 'ready', job
        download = client.get(job['download_url'])
        assert download.status_code == 200
        assert download.headers['content-type'] == 'application/zip'
        with ZipFile(BytesIO(download.content)) as archive:
            assert archive.testzip() is None
            images = [name for name in archive.namelist() if name.endswith('.jpg')]
            assert len(images) == 1 and images[0].endswith('/one/good.jpg')
            assert archive.read(images[0]) == (media / 'good.jpg').read_bytes()
            book = load_workbook(BytesIO(archive.read('审核结果.xlsx')))
            assert book['审核结果'].max_column == 16
            assert book['审核结果'].cell(2, 8).value == '=literal'
            assert book['图片对应表'].cell(2, 5).value == images[0]
            assert '待审核帖子：1' in archive.read('导出说明.txt').decode('utf-8-sig')
            book.close()
        assert client.delete('/api/data/review/export/jobs/' + job['id']).status_code == 200
        assert client.get(job['download_url']).status_code in (404, 409)
    assert (path.read_bytes(), read_state(tmp_path, path)) == before


@pytest.mark.asyncio
async def test_interrupted_asgi_download_releases_package_for_removal(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    path = tmp_path / 'weibo/posts.json'; path.parent.mkdir()
    path.write_text('[{"note_id":"empty"}]', encoding='utf-8')
    update_annotation(tmp_path, path, 'empty:', 'post', 'qualified')
    app = FastAPI(); app.include_router(router.router, prefix='/api')
    with TestClient(app) as client:
        job = ready(client, client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json())
        scope = {'type':'http', 'asgi':{'version':'3.0'}, 'http_version':'1.1', 'method':'GET',
                 'scheme':'http', 'path':job['download_url'], 'raw_path':job['download_url'].encode(),
                 'query_string':b'', 'root_path':'', 'headers':[], 'server':('testserver',80), 'client':('testclient',1)}
        async def receive():
            return {'type':'http.request', 'body':b'', 'more_body':False}
        async def fail_send(message):
            if message['type'] == 'http.response.body':
                raise OSError('simulated client disconnect')
        with pytest.raises(OSError, match='simulated client disconnect'):
            await app(scope, receive, fail_send)
        assert client.delete('/api/data/review/export/jobs/' + job['id']).status_code == 200
        assert not list((tmp_path / '.review-exports').glob('*.zip'))


def test_unicode_original_filename_exports_with_portable_name(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    name = '长' * 78 + '.jpg'
    path = tmp_path / 'weibo/posts.json'; path.parent.mkdir()
    path.write_text(json.dumps([{'note_id':'unicode','image_names':name}],ensure_ascii=False),encoding='utf-8')
    media = tmp_path / 'wb/media/unicode'; media.mkdir(parents=True)
    Image.new('RGB',(20,30)).save(media/name)
    update_annotation(tmp_path,path,'unicode:','image','qualified',name)
    app=FastAPI();app.include_router(router.router,prefix='/api')
    with TestClient(app) as client:
        job=ready(client,client.post('/api/data/review/export/jobs',json={'file':'weibo/posts.json'}).json())
        assert job['status']=='ready',job
        with ZipFile(BytesIO(client.get(job['download_url']).content)) as archive:
            images=[n for n in archive.namelist() if n.endswith('.jpg')]
            assert len(images)==1 and len(images[0].split('/')[-1].encode('utf-8'))<=180
            archive.extractall(tmp_path/'unpacked')
            assert (tmp_path/'unpacked'/images[0]).read_bytes()==(media/name).read_bytes()
        client.delete('/api/data/review/export/jobs/'+job['id'])


def test_file_metadata_maps_only_main_files_in_saved_task(tmp_path, monkeypatch):
    import config
    from collection_tasks import create_task
    monkeypatch.setattr(config,'SAVE_DATA_PATH',str(tmp_path))
    folder=tmp_path/'weibo';folder.mkdir()
    (folder/'first.xlsx').write_bytes(b'fixture')
    (folder/'first_semantic.xlsx').write_bytes(b'fixture')
    (folder/'orphan_semantic.xlsx').write_bytes(b'fixture')
    create_task(tmp_path,{}, {},anchor='weibo/first.xlsx')
    app=FastAPI();app.include_router(router.router,prefix='/api')
    with TestClient(app) as client:
        files={item['name']:item for item in client.get('/api/data/files').json()['files']}
        assert files['first_semantic.xlsx']['review_main']=='weibo/first.xlsx'
        assert files['orphan_semantic.xlsx']['review_main'] is None


def test_storage_full_reports_failure_without_changing_review(tmp_path, monkeypatch):
    import io
    import errno
    import config
    monkeypatch.setattr(config,'SAVE_DATA_PATH',str(tmp_path))
    path=tmp_path/'weibo/posts.json';path.parent.mkdir()
    path.write_text('[{"note_id":"empty"}]',encoding='utf-8')
    update_annotation(tmp_path,path,'empty:','post','qualified')
    before=path.read_bytes(),read_state(tmp_path,path)
    original_open=io.open
    def disk_full(file,mode='r',*args,**kwargs):
        if str(file).endswith('.zip') and 'w' in mode:
            raise OSError(errno.ENOSPC,'simulated full export storage')
        return original_open(file,mode,*args,**kwargs)
    monkeypatch.setattr(io,'open',disk_full)
    app=FastAPI();app.include_router(router.router,prefix='/api')
    with TestClient(app) as client:
        job=ready(client,client.post('/api/data/review/export/jobs',json={'file':'weibo/posts.json'}).json())
        assert job['status']=='failed' and job['error']=='review_export_failed'
        assert not list((tmp_path/'.review-exports').glob('*.zip'))
    assert (path.read_bytes(),read_state(tmp_path,path))==before


def test_missing_approved_image_and_duplicate_references_are_explicit(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    path = tmp_path / 'weibo/posts.json'
    path.parent.mkdir()
    notes = [dict(note_id='same', source_keyword=word, image_names='good.jpg,missing.jpg') for word in ('可爱 公园', '萌 公园')]
    notes.append(dict(note_id='empty', image_names='none.jpg'))
    path.write_text(json.dumps(notes, ensure_ascii=False), encoding='utf-8')
    folder = tmp_path / 'wb/media/same'
    folder.mkdir(parents=True)
    Image.new('RGB', (30, 30)).save(folder / 'good.jpg')
    for word in ('可爱 公园', '萌 公园'):
        update_annotation(tmp_path, path, 'same:' + word, 'image', 'qualified', 'good.jpg')
        update_annotation(tmp_path, path, 'same:' + word, 'image', 'qualified', 'missing.jpg')
    update_annotation(tmp_path, path, 'empty:', 'post', 'qualified')
    before = read_state(tmp_path, path)
    app = FastAPI(); app.include_router(router.router, prefix='/api')
    with TestClient(app) as client:
        job = ready(client, client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json())
        assert job['status'] == 'ready', job
        assert job['counts'] == {'retained': 3, 'images': 1, 'pending': 0, 'missing': 2, 'no_images': 1}
        with ZipFile(BytesIO(client.get(job['download_url']).content)) as archive:
            assert len([n for n in archive.namelist() if n.endswith('.jpg')]) == 1
            book = load_workbook(BytesIO(archive.read('审核结果.xlsx')))
            rows = list(book['图片对应表'].values)[1:]
            assert rows[0][4] == rows[2][4]
            assert rows[1][4] is None and rows[1][5] == '未下载'
            assert rows[-1][5] == '无可导出合格图片'
            book.close()
        assert read_state(tmp_path, path) == before
        client.delete('/api/data/review/export/jobs/' + job['id'])


def test_empty_traversal_download_and_root_isolation(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    path = tmp_path / 'weibo/posts.json'; path.parent.mkdir()
    path.write_text('[{"note_id":"empty"}]', encoding='utf-8')
    app = FastAPI(); app.include_router(router.router, prefix='/api')
    with TestClient(app) as client:
        assert client.post('/api/data/review/export/jobs', json={'file':'../outside.json'}).status_code == 403
        assert client.get('/api/data/review/export/jobs/not-a-job').status_code == 404
        job = ready(client, client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json())
        assert job['status'] == 'failed' and job['error'] == 'review_export_empty'
        update_annotation(tmp_path, path, 'empty:', 'post', 'qualified')
        job = ready(client, client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json())
        response = client.get(job['download_url'])
        assert response.status_code == 200 and response.content.startswith(b'PK\x03\x04')
        monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path / 'another-root'))
        assert client.get(job['download_url']).status_code == 404
        monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
        assert client.delete('/api/data/review/export/jobs/' + job['id']).status_code == 200


def test_cancel_capacity_and_expiry_preserve_original_files(tmp_path, monkeypatch):
    import config
    from api.services import review_export_jobs as jobs_module
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    path = tmp_path / 'weibo/posts.json'; path.parent.mkdir()
    path.write_text('[{"note_id":"empty"}]', encoding='utf-8')
    update_annotation(tmp_path, path, 'empty:', 'post', 'qualified')
    before = path.read_bytes(), read_state(tmp_path, path)
    app = FastAPI(); app.include_router(router.router, prefix='/api')
    with TestClient(app) as client:
        job = client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json()
        assert client.delete('/api/data/review/export/jobs/' + job['id']).status_code == 200
        cancelled = ready(client, job)
        assert cancelled['status'] == 'cancelled'
        assert not list((tmp_path / '.review-exports').glob('*.zip'))
        jobs = [ready(client, client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'}).json()) for _ in range(3)]
        assert all(job['status'] == 'ready' for job in jobs)
        response = client.post('/api/data/review/export/jobs', json={'file':'weibo/posts.json'})
        assert response.status_code == 409 and response.json()['detail'] == 'review_export_capacity'
        future = time.time() + 3700
        monkeypatch.setattr(jobs_module.time, 'time', lambda: future)
        assert client.get('/api/data/review/export/jobs/' + jobs[0]['id']).status_code == 404
        assert not list((tmp_path / '.review-exports').glob('*.zip'))
    assert (path.read_bytes(), read_state(tmp_path, path)) == before
