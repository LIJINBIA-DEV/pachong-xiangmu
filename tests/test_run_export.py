"""Regression for the 10 matched + 33 review run and historical export leakage."""
import httpx
import openpyxl
import pytest

from semantic_filter import StreamingSemanticFilter
from tests.semantic_responses import grounded_response


@pytest.mark.asyncio
async def test_export_only_current_run_even_when_checkpoint_has_more_matches(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        old = StreamingSemanticFilter(tmp_path / 'old.xlsx', [('可爱', '公园')], 'fake', client=client)
        for i in range(16):
            await old.add_note({'note_id': str(i), 'content': '这个公园好可爱', 'source_keyword': '可爱 公园', "image_names": "001.jpg"})
            await old.finalize_note(str(i))
        await old.close()
        current = StreamingSemanticFilter(tmp_path / 'new.xlsx', [('可爱', '公园')], 'fake', client=client)
        for i in range(10):
            await current.add_note({'note_id': str(i), 'content': '这个公园好可爱', 'source_keyword': '可爱 公园', "image_names": "001.jpg"})
            await current.finalize_note(str(i))
        await current.close()
        assert len(current.states) == 16  # Cache retained for future resumptions.
        assert current.reviewer.calls == 0
    book = openpyxl.load_workbook(current.output_path)
    assert book['Matched'].max_row == 11
    assert dict(book['Summary'].values)['matched'] == 10
    assert dict(book['Summary'].values)['api_calls'] == 0


@pytest.mark.asyncio
@pytest.mark.parametrize('matched_count', [10, 0])
async def test_primary_excel_separates_33_reviews_and_their_comments(tmp_path, monkeypatch, matched_count):
    import config
    from store.excel_store_base import ExcelStoreBase
    from types import SimpleNamespace
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER', raising=False)
    store = ExcelStoreBase('weibo', 'search')
    stream = SimpleNamespace(states={}, note_keys={})
    store.semantic_stream = stream
    for i in range(33 + matched_count):
        nid = str(i)
        key = (nid, '可爱 公园')
        note = {'note_id': nid, 'content': '测试场景', 'source_keyword': key[1], "image_names": "001.jpg"}
        stream.states[key] = SimpleNamespace(status='matched' if i < matched_count else 'review', note=note)
        stream.note_keys[nid] = {key}
        store._semantic_pending_notes[key] = note
        store._semantic_pending_comments[nid] = [{'comment_id': f'c{i}', 'note_id': nid, 'content': '评论'}]
        store._write_retained_semantic_note(nid)
    store.flush()
    book = openpyxl.load_workbook(store.filename)
    assert book.sheetnames[0] == 'Contents'
    assert book['Contents'].max_row == matched_count + 1
    assert book['Review'].max_row == 34
    assert book['ReviewComments'].max_row == 34
    if matched_count:
        assert book['Comments'].max_row == matched_count + 1
