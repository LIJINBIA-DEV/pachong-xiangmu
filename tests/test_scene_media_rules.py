"""V4 animal boundaries and image presence gates; model responses are controlled."""
import hashlib
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import openpyxl
import pytest

from semantic_filter import StreamingSemanticFilter, CHECKPOINT_VERSION, MIN_KEEP_CONFIDENCE, MODEL
from semantic_quality import image_requirement_problem
from tests.semantic_responses import grounded_response


@pytest.mark.parametrize(('note', 'expected'), [
    ({'image_names': '001.jpg,002.png'}, None),
    ({'image_names': 'cover.webp'}, None),
    ({'image_list': 'https://example.com/picture'}, None),
    ({'image_list': "['https://example.com/001.jpg']"}, None),
    ({'images': [{'url_default': 'https://example.com/image.jpg'}]}, None),
    ({'image_names': ''}, 'no_images'),
    ({'image_names': 'video.mp4'}, 'no_images'),
    ({'image_list': []}, 'no_images'),
    ({'image_list': None}, 'no_images'),
    ({'image_list': '[invalid'}, 'no_images'),
    ({'image_list': 'javascript:alert(1)'}, 'no_images'),
    ({'image_list': 'https://example.com/video.mp4'}, 'no_images'),
    ({}, 'image_information_unavailable'),
])
def test_image_metadata_gate(note, expected):
    assert image_requirement_problem(note) == expected


@pytest.mark.asyncio
@pytest.mark.parametrize(('content', 'kind', 'choice', 'expected'), [
    ('公园里的野生松鼠真可爱', 'wild_animal', 'keep', 'matched'),
    ('公园池塘的野生鸭子真可爱', 'wild_animal', 'keep', 'matched'),
    ('公园动物区饲养的小鹿真可爱', 'park_animal', 'keep', 'matched'),
    ('带我家可爱的小狗逛公园', 'pet', 'reject', 'excluded'),
    ('别人的宠物狗在公园好可爱', 'pet', 'reject', 'excluded'),
    ('公园里的流浪猫真可爱', 'animal', 'review', 'matched'),
    ('公园里的小猫真可爱', 'animal', 'review', 'matched'),
    ('公园里的小猫真可爱', 'wild_animal', 'keep', 'matched'),
    ('公园里的流浪猫真可爱', 'wild_animal', 'keep', 'matched'),
    # Explicit pet evidence must not be accepted through review promotion.
    ('我家宠物狗在公园真可爱', 'wild_animal', 'keep', 'excluded'),
])
async def test_animal_boundaries_with_picture(tmp_path, content, kind, choice, expected):
    def respond(request):
        return grounded_response(request, choice, kind=kind, relation='scene_component')
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client)
        await stream.add_note({'note_id': '1', 'source_keyword': '可爱 公园', 'content': content, 'image_names': '001.jpg'})
        await stream.finalize_note('1')
        await stream.close()
    assert stream.states[('1', '可爱 公园')].status == expected


@pytest.mark.asyncio
async def test_nearby_wildlife_does_not_pass_specific_element(tmp_path):
    def respond(request):
        return grounded_response(request, kind='wild_animal', relation='target_itself')
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '垃圾桶')], 'fake', client=client)
        await stream.add_note({'note_id': '1', 'source_keyword': '可爱 垃圾桶', 'content': '垃圾桶旁边的松鼠真可爱', 'image_names': '001.jpg'})
        await stream.finalize_note('1')
        await stream.close()
    state = stream.states[('1', '可爱 垃圾桶')]
    assert state.status == 'matched' and state.quality_reason == 'animal_outside_place_scope'
    assert state.note['selection_origin'] == 'review_policy'


@pytest.mark.asyncio
@pytest.mark.parametrize('repair', [True, False])
async def test_explicit_animal_rejection_is_final_without_repair(tmp_path, repair):
    sent = []
    def respond(request):
        sent.append(json.loads(request.content))
        if len(sent) == 1 or not repair:
            return grounded_response(request, 'reject', kind='animal', relation='unrelated')
        return grounded_response(request, kind='wild_animal', relation='scene_component')
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client)
        await stream.add_note({'note_id': '1', 'source_keyword': '可爱 公园', 'content': '公园里的野生松鼠真可爱', 'image_names': '001.jpg'})
        await stream.finalize_note('1')
        await stream.close()
    assert len(sent) == 1
    assert stream.states[('1', '可爱 公园')].status == 'excluded'


@pytest.mark.asyncio
async def test_no_images_excluded_before_model_and_unknown_images_retained_for_manual_review(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'run.xlsx', [('可爱', '公园')], 'fake', client=client)
        await stream.add_note({'note_id': 'empty', 'source_keyword': '可爱 公园', 'content': '这个公园好可爱', 'image_names': ''})
        assert stream.states[('empty', '可爱 公园')].source == 'no_images'
        assert stream.reviewer.calls == 0
        await stream.finalize_note('empty', comments_available=False)
        assert stream.reviewer.calls == 0
        await stream.add_note({'note_id': 'unknown', 'source_keyword': '可爱 公园', 'content': '这个公园好可爱'})
        await stream.finalize_note('unknown')
        await stream.close()
    book = openpyxl.load_workbook(stream.output_path, read_only=True)
    try:
        assert book['Matched'].max_row == 2
        assert book['Review'].max_row == 1
        assert dict(book['Summary'].values)['excluded'] == 1
        assert dict(book['Summary'].values)['matched'] == 1
    finally:
        book.close()
    assert stream.states[('unknown', '可爱 公园')].source == 'image_information_unavailable'


@pytest.mark.asyncio
async def test_v3_checkpoint_not_reused_or_overwritten(tmp_path):
    identity = {'version': CHECKPOINT_VERSION, 'rule': 'scene-evidence-v3', 'pairs': [('可爱', '公园')],
                'model': MODEL, 'scope': 'live', 'confidence_threshold': MIN_KEEP_CONFIDENCE}
    digest = hashlib.sha256(json.dumps(identity, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()[:12]
    old = tmp_path / f'semantic_checkpoint_{digest}.jsonl'
    original = json.dumps({'note': {'note_id': '1', 'source_keyword': '可爱 公园', 'content': '公园里的松鼠真可爱'}, 'status': 'excluded'}).encode()
    old.write_bytes(original)
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'new.xlsx', [('可爱', '公园')], 'fake', client=client)
        assert stream.checkpoint_path != old and not stream.states
        await stream.close()
    assert old.read_bytes() == original


@pytest.mark.asyncio
async def test_live_search_skips_pictureless_comments_and_counts_only_pictured_match(tmp_path, monkeypatch):
    from media_platform.weibo import core
    from store.excel_store_base import ExcelStoreBase
    import config
    monkeypatch.setattr(config, 'ENABLE_GET_COMMENTS', False)
    for name, value in {'KEYWORDS': '可爱 公园', 'START_PAGE': 1, 'CRAWLER_MAX_NOTES_COUNT': 1,
                        'CRAWLER_MAX_SLEEP_SEC': 0, 'SAVE_DATA_OPTION': 'excel', 'SAVE_DATA_PATH': str(tmp_path)}.items():
        monkeypatch.setattr(config, name, value)
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER', raising=False)
    store = ExcelStoreBase('weibo', 'search')
    monkeypatch.setenv('MEDIACRAWLER_SEMANTIC_FILTER', '1')
    monkeypatch.setattr(ExcelStoreBase, '_instances', {'weibo_search': store})
    monkeypatch.setattr(core.weibo_store.WeibostoreFactory, 'create_store', lambda: store)
    monkeypatch.setattr(core, 'filter_search_result_card', lambda cards: cards)
    crawler = core.WeiboCrawler()
    cards = [{'mblog': {'id': nid, 'text': '这个公园好可爱', 'created_at': 'Thu Oct 08 09:00:00 +0800 2026',
                         'pics': [] if nid == 'empty' else [{'url': 'https://wx4.sinaimg.cn/large/test.jpg'}]}} for nid in ('empty', 'pictured')]
    crawler.wb_client = SimpleNamespace(get_note_by_keyword=AsyncMock(return_value={'cards': cards}))
    crawler.batch_get_notes_full_text = AsyncMock(side_effect=lambda notes: notes)
    async def finalize(ids):
        for nid in ids:
            await ExcelStoreBase.finalize_semantic_note('weibo', nid)
    crawler.batch_get_notes_comments = AsyncMock(side_effect=finalize)
    crawler.download_media = AsyncMock()
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        store.semantic_stream = StreamingSemanticFilter(store.filename, [('可爱', '公园')], 'fake', client=client)
        await crawler.search()
        await store.semantic_stream.close()
    store.flush()
    # The pictured post qualifies from its body; neither post needs comments.
    assert crawler.batch_get_notes_comments.await_count == 0
    assert [call.args[0]['id'] for call in crawler.download_media.await_args_list] == ['pictured']
    book = openpyxl.load_workbook(store.filename, read_only=True)
    try:
        rows = list(book['Contents'].values)
        assert len(rows) == 2 and rows[1][0] == 'pictured'
    finally:
        book.close()
