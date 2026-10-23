"""Regression checks for rule routing, evidence, recovery and offline parity.

Mock responses exercise the integration, not the model's real-world accuracy.
"""

import hashlib
import json

import httpx
import openpyxl
import pytest

from semantic_filter import JevReviewer, StreamingSemanticFilter, filter_checkpoint, filter_workbook
from semantic_rules import RULE_VERSION, relation_instructions, target_scope
from tests.semantic_responses import grounded_response


@pytest.mark.parametrize("target,scope", [
    ("公园", "place"), ("咖啡馆", "place"), ("雕塑", "element"),
    ("垃圾桶", "element"), ("这里", "contextual_place"), ("citywalk", "contextual_place"),
    ("自定义对象", "specific_target"),
])
def test_keyword_scope(target, scope):
    assert target_scope(target) == scope


def answer(request, *, choice="keep", kind="landscape", relation="scene_component"):
    return grounded_response(request, choice, kind=kind, relation=relation)


@pytest.mark.asyncio
@pytest.mark.parametrize("target,kind,relation,expected", [
    ("公园", "landscape", "scene_component", "keep"),
    ("公园", "building", "scene_component", "keep"),
    ("咖啡馆", "decoration", "scene_component", "keep"),
    ("公园", "animal", "scene_component", "keep"),
    ("公园", "person", "scene_component", "review"),
    ("公园", "food", "scene_component", "review"),
    ("公园", "unrelated_item", "unrelated", "review"),
    ("垃圾桶", "facility", "scene_component", "review"),
    ("垃圾桶", "facility", "element_component", "keep"),
])
async def test_scope_prompt_and_conflicting_model_answers(target, kind, relation, expected):
    sent = []
    def respond(request):
        sent.append(json.loads(request.content))
        return answer(request, kind=kind, relation=relation)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        reviewer = JevReviewer("fake", client)
        decision = await reviewer.decide(("可爱", target), f"{target}里的猫形雕塑真可爱")
    assert decision.choice == expected
    assert sent[0]["questions"]["relation"]["instructions"] == relation_instructions(target)
    assert "真实动物" in sent[0]["questions"]["relation"]["instructions"]
    assert "不能当作真实动物排除" in sent[0]["questions"]["relation"]["instructions"]


@pytest.mark.asyncio
async def test_live_and_offline_use_identical_context_and_save_evidence(tmp_path):
    sent = []
    def respond(request):
        sent.append(json.loads(request.content))
        return answer(request)
    note = {"note_id": "1", "content": "公园里的猫形雕塑真可爱", "source_keyword": "可爱 公园", "semantic_comments_available": True, "image_names": "001.jpg"}
    comment = {"note_id": "1", "comment_id": "a", "content": "门口的小房子也很可爱"}
    raw = tmp_path / "raw.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Contents"
    book.active.append(list(note))
    book.active.append(list(note.values()))
    sheet = book.create_sheet("Comments")
    sheet.append(list(comment))
    sheet.append(list(comment.values()))
    book.save(raw)
    original = raw.read_bytes()
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        live = StreamingSemanticFilter(tmp_path / "live.xlsx", [("可爱", "公园")], "fake", client=client)
        await live.add_note(note)
        await live.add_comment(comment)
        await live.finalize_note("1")
        await live.close()
        output, counts = await filter_workbook(raw, [("可爱", "公园")], "fake", client)
    assert sent[:2] == sent[2:]
    assert counts["matched"] == 1
    assert raw.read_bytes() == original
    result = openpyxl.load_workbook(output)
    rows = list(result["Matched"].values)
    record = dict(zip(rows[0], rows[1]))
    assert "猫形雕塑真可爱" in record["semantic_evidence"]
    assert record["semantic_object_type"].startswith("雕塑")
    assert record["semantic_rule_version"] == RULE_VERSION


@pytest.mark.asyncio
async def test_pending_survives_failure_and_resumes_with_saved_comments(tmp_path, monkeypatch):
    import semantic_filter
    async def no_sleep(_):
        pass
    monkeypatch.setattr(semantic_filter.asyncio, "sleep", no_sleep)
    def fail(request):
        raise httpx.ConnectError("offline", request=request)
    async with httpx.AsyncClient(transport=httpx.MockTransport(fail)) as client:
        first = StreamingSemanticFilter(tmp_path / "first.xlsx", [("可爱", "公园")], "fake", client=client)
        await first.add_note({"note_id": "1", "content": "公园里的猫形雕塑真可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await first.add_comment({"note_id": "1", "comment_id": "a", "content": "同意"})
        await first.finalize_note("1")
        assert first.states[("1", "可爱 公园")].status == "pending"
        # More candidates can still be saved during the provider cooldown.
        await first.add_note({"note_id": "2", "content": "可爱的公园", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await first.finalize_note("2")
        assert first.reviewer.calls == 4
        await first.close()
    async with httpx.AsyncClient(transport=httpx.MockTransport(answer)) as client:
        second = StreamingSemanticFilter(tmp_path / "second.xlsx", [("可爱", "公园")], "fake", client=client)
        assert second.states[("1", "可爱 公园")].comments[0]["comment_id"] == "a"
        await second.finalize_note("1")
        await second.close()
    assert second.states[("1", "可爱 公园")].status == "matched"
    assert second.states[("1", "可爱 公园")].comments_available is True


@pytest.mark.asyncio
async def test_v1_cache_is_not_reused_but_can_be_explicitly_rechecked(tmp_path):
    pairs = [("可爱", "公园")]
    old_identity = json.dumps({"version": 1, "pairs": pairs, "model": "jev-1.13.0"}, ensure_ascii=False, sort_keys=True)
    digest = hashlib.sha256(old_identity.encode()).hexdigest()[:12]
    old = tmp_path / f"semantic_checkpoint_{digest}.jsonl"
    row = {"note": {"note_id": "1", "content": "公园里的猫形雕塑真可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"}, "comments": [], "status": "excluded", "source": "semantic_reject"}
    old.write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")
    original = old.read_bytes()
    async with httpx.AsyncClient(transport=httpx.MockTransport(answer)) as client:
        live = StreamingSemanticFilter(tmp_path / "live.xlsx", pairs, "fake", client=client)
        assert not live.states
        assert live.checkpoint_path != old
        await live.close()
        output, counts = await filter_checkpoint(old, pairs, "fake", client)
    assert counts["matched"] == 1
    assert old.read_bytes() == original
    assert output.exists()


@pytest.mark.asyncio
async def test_incomplete_comments_are_retried_on_new_crawl(tmp_path):
    def reject(request):
        return answer(request, choice="reject", kind="animal", relation="unrelated")
    note = {"note_id": "1", "content": "公园散步", "source_keyword": "可爱 公园", "image_names": "001.jpg"}
    async with httpx.AsyncClient(transport=httpx.MockTransport(reject)) as client:
        first = StreamingSemanticFilter(tmp_path / "one.xlsx", [("可爱", "公园")], "fake", client=client)
        await first.add_note(note)
        await first.finalize_note("1", False)
        await first.close()
        second = StreamingSemanticFilter(tmp_path / "two.xlsx", [("可爱", "公园")], "fake", client=client)
        await second.add_note(note)
        assert second.states[("1", "可爱 公园")].status is None
        await second.add_comment({"note_id": "1", "comment_id": "a", "content": "公园的雕塑也可爱"})
        assert len(second.states[("1", "可爱 公园")].comments) == 1
        await second.close()


@pytest.mark.asyncio
async def test_pending_limit_stops_safely_and_shutdown_never_promotes_pending(tmp_path, monkeypatch):
    import semantic_filter
    from semantic_filter import SemanticFilterError
    monkeypatch.setattr(semantic_filter, "MAX_PENDING_NOTES", 2)
    async def no_sleep(_):
        pass
    monkeypatch.setattr(semantic_filter.asyncio, "sleep", no_sleep)
    calls = []
    def fail(request):
        calls.append(request)
        return httpx.Response(503)
    async with httpx.AsyncClient(transport=httpx.MockTransport(fail)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        for note_id in ("1", "2"):
            await stream.add_note({"note_id": note_id, "content": "公园里的雕塑好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
            if note_id == "1":
                await stream.finalize_note(note_id)
            else:
                with pytest.raises(SemanticFilterError, match="2 pending posts saved"):
                    await stream.finalize_note(note_id)
        # Simulate recovery at shutdown: shutdown still must not bypass quotas.
        stream.reviewer.failure = None
        await stream.close()
    assert len(calls) == 4
    assert all(state.status == "pending" for state in stream.states.values())
    rows = [json.loads(line) for line in stream.checkpoint_path.read_text(encoding="utf-8").splitlines()]
    assert {row["note"]["note_id"] for row in rows} == {"1", "2"}


@pytest.mark.asyncio
async def test_comment_keyword_gate_does_not_join_separate_comments(tmp_path):
    def unexpected(request):
        raise AssertionError("No single field contains the pair; do not call the model")
    async with httpx.AsyncClient(transport=httpx.MockTransport(unexpected)) as client:
        stream = StreamingSemanticFilter(tmp_path / "run.xlsx", [("可爱", "公园")], "fake", client=client)
        await stream.add_note({"note_id": "1", "content": "今天出门", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await stream.add_comment({"note_id": "1", "comment_id": "a", "content": "好可爱"})
        await stream.add_comment({"note_id": "1", "comment_id": "b", "content": "这是公园"})
        await stream.finalize_note("1")
        await stream.close()
    assert stream.states[("1", "可爱 公园")].source == "keyword_miss"
