import json
import importlib
from types import SimpleNamespace

import httpx
import openpyxl
import pytest

from semantic_filter import LiveNote, StreamingSemanticFilter, _note_context_chunks, check_jev_connection, filter_workbook
from api.schemas.crawler import CrawlerStartRequest
from api.services.crawler_manager import CrawlerManager
from tests.semantic_responses import grounded_response


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("status", "expected_reason"),
    [(200, "usable"), (403, "authentication_rejected"), (520, "provider_error")],
)
async def test_semantic_key_check_reports_real_api_status_without_exposing_key(status, expected_reason):
    def respond(request):
        assert request.headers["Authorization"] == "Bearer test-secret"
        if status != 200:
            return httpx.Response(status)
        return grounded_response(request)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        result = await check_jev_connection("test-secret", client=client)

    assert result["reason"] == expected_reason
    assert result["ok"] is (status == 200)
    assert "test-secret" not in repr(result)


@pytest.mark.asyncio
async def test_filters_same_field_and_individual_comment_without_changing_raw(tmp_path):
    raw_path = tmp_path / "wb_search_test.xlsx"
    book = openpyxl.Workbook()
    notes = book.active
    notes.title = "Contents"
    notes.append(["note_id", "content", "source_keyword", "semantic_comments_available", "image_names"])
    notes.append(["1", "这个公园好可爱", "可爱 公园", True, "001.jpg"])
    notes.append(["2", "公园里的小狗好可爱", "可爱 公园", True, "001.jpg"])
    notes.append(["3", "公园在这里", "可爱 公园", True, "001.jpg"])
    notes.append(["4", "今天去了公园", "可爱 公园", True, "001.jpg"])
    notes.append(["5", "可爱广场", "可爱 公园", True, "001.jpg"])
    comments = book.create_sheet("Comments")
    comments.append(["comment_id", "note_id", "content"])
    comments.append(["31", "3", "真可爱"])
    comments.append(["32", "3", "这个公园不错"])
    comments.append(["41", "4", "这个公园好可爱"])
    book.save(raw_path)

    sent = []

    def respond(request):
        payload = json.loads(request.content)
        sent.append(payload)
        text = payload["state"]["text"]
        choice = "reject" if "小狗" in text else "keep"
        return grounded_response(request, choice, source="comment:41" if "一级评论 41" in text else None)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        output_path, counts = await filter_workbook(raw_path, [("可爱", "公园")], "test-key", client)

    assert counts["matched"] == 2
    assert counts["excluded"] == 3
    assert counts["review"] == 0
    assert counts["api_calls"] == 5  # Two keeps verified; one reject.
    assert all(item["state"]["keyword2"] == "公园" for item in sent)
    assert "今天去了公园" in sent[-1]["state"]["text"]
    assert "一级评论 41" in sent[-1]["state"]["text"]

    filtered = openpyxl.load_workbook(output_path)
    assert [row[0] for row in list(filtered["Matched"].values)[1:]] == ["1", "4"]
    assert "Excluded" not in filtered.sheetnames
    assert dict(filtered["Summary"].values)["excluded"] == 3
    assert list(filtered["MatchedComments"].values)[1][0] == "41"
    assert openpyxl.load_workbook(raw_path)["Contents"].max_row == 6


@pytest.mark.asyncio
async def test_api_failure_stops_filtering_without_marking_posts_for_review(tmp_path):
    raw_path = tmp_path / "xhs_search_test.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Contents"
    sheet.append(["note_id", "title", "desc", "source_keyword"])
    sheet.append(["1", "可爱公园", "", "可爱 公园"])
    sheet.append(["2", "可爱公园", "", "可爱 公园"])
    book.save(raw_path)
    sent = []

    def fail(request):
        sent.append(request)
        return httpx.Response(401, json={"error": "invalid key"})

    async with httpx.AsyncClient(transport=httpx.MockTransport(fail)) as client:
        with pytest.raises(RuntimeError, match="Jev API failed"):
            await filter_workbook(raw_path, [("可爱", "公园")], "test-key", client)

    assert len(sent) == 1
    result = openpyxl.load_workbook(tmp_path / "xhs_search_test_semantic.xlsx")
    assert dict(result["Summary"].values)["pending"] == 2
    assert result["Matched"].max_row == 1
    assert list(tmp_path.glob("semantic_checkpoint_*.jsonl"))


@pytest.mark.asyncio
async def test_streaming_api_failure_does_not_retain_unjudged_post(tmp_path, monkeypatch):
    import semantic_filter
    async def no_sleep(delay):
        pass
    monkeypatch.setattr(semantic_filter.asyncio, "sleep", no_sleep)
    def fail(request):
        return httpx.Response(520)

    async with httpx.AsyncClient(transport=httpx.MockTransport(fail)) as client:
        reviewer = StreamingSemanticFilter(
            tmp_path / "weibo_search_test.xlsx", [("可爱", "公园")], "test-key", client=client,
        )
        await reviewer.add_note({
            "note_id": "1", "content": "公园里的小狗好可爱", "source_keyword": "可爱 公园",
                                                                      "image_names": "001.jpg",
        })
        await reviewer.finalize_note("1")
        await reviewer.close()

    assert reviewer.states[("1", "可爱 公园")].status == "pending"
    result = openpyxl.load_workbook(reviewer.output_path)
    assert result["Matched"].max_row == 1
    assert result["Review"].max_row == 1
    assert dict(result["Summary"].values)["pending"] == 1
    assert dict(result["Summary"].values)["excluded"] == 0
    assert "520" in dict(result["Summary"].values)["api_failure"]
    assert reviewer.checkpoint_path.exists()
    saved = json.loads(reviewer.checkpoint_path.read_text(encoding="utf-8").splitlines()[-1])
    assert saved["status"] == "pending"
    assert saved["note"]["content"] == "公园里的小狗好可爱"


@pytest.mark.asyncio
async def test_transient_jev_connection_failure_retries_before_saving_decision(tmp_path, monkeypatch):
    import semantic_filter

    delays = []

    async def fake_sleep(delay):
        delays.append(delay)

    monkeypatch.setattr(semantic_filter.asyncio, "sleep", fake_sleep)
    attempts = []

    def respond(request):
        attempts.append(request)
        if len(attempts) == 1:
            raise httpx.ConnectError("proxy disconnected", request=request)
        return grounded_response(request)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        reviewer = StreamingSemanticFilter(
            tmp_path / "weibo_search_test.xlsx", [("可爱", "公园")], "test-key", client=client,
        )
        await reviewer.add_note({"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await reviewer.finalize_note("1")
        await reviewer.close()

    assert len(attempts) == 3
    assert delays == [1]
    assert reviewer.states[("1", "可爱 公园")].status == "matched"
    assert dict(openpyxl.load_workbook(reviewer.output_path)["Summary"].values)["api_calls"] == 3
    rows = [json.loads(line) for line in reviewer.checkpoint_path.read_text(encoding="utf-8").splitlines()]
    assert [row["status"] for row in rows] == ["pending", "matched"]


@pytest.mark.asyncio
async def test_semantic_start_requires_key_before_crawling(monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    manager = CrawlerManager()
    request = CrawlerStartRequest(
        platform="wb", crawler_type="search", save_option="excel",
        keywords="可爱 公园", keyword_pairs=[("可爱", "公园")], semantic_filter=True,
    )
    with pytest.raises(ValueError, match="TYPESAFE_API_KEY"):
        await manager.start(request)
    assert manager.process is None


@pytest.mark.asyncio
async def test_web_key_and_model_reach_child_environment_without_entering_logs(monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    captured = {}

    class FakeProcess:
        def poll(self):
            return None

    def fake_popen(command, **kwargs):
        captured["command"] = command
        captured["env"] = kwargs["env"]
        return FakeProcess()

    manager_module = importlib.import_module("api.services.crawler_manager")
    monkeypatch.setattr(manager_module.subprocess, "Popen", fake_popen)
    manager = CrawlerManager()
    request = CrawlerStartRequest(
        platform="wb", crawler_type="search", save_option="excel",
        keywords="可爱 公园", keyword_pairs=[("可爱", "公园")],
        semantic_filter=True, semantic_model="jev-1.13.0",
        semantic_api_key="web-test-secret",
    )
    assert await manager.start(request)
    assert captured["env"]["TYPESAFE_API_KEY"] == "web-test-secret"
    assert captured["env"]["MEDIACRAWLER_SEMANTIC_MODEL"] == "jev-1.13.0"
    assert "web-test-secret" not in " ".join(captured["command"])
    assert "web-test-secret" not in " ".join(entry.message for entry in manager.logs)
    assert "web-test-secret" not in repr(request)
    assert manager.current_config.semantic_api_key is None
    manager._read_task.cancel()


@pytest.mark.asyncio
async def test_semantic_review_waits_for_all_comments_and_uses_post_context(tmp_path):
    sent = []

    def respond(request):
        text = json.loads(request.content)["state"]["text"]
        sent.append(text)
        choice = "keep" if "这个公园好可爱" in text else "reject"
        evidence = "comment_1" if "一级评论 2a" in text else "post"
        return grounded_response(request, choice, source="comment:2a" if "一级评论 2a" in text else None)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        reviewer = StreamingSemanticFilter(
            tmp_path / "weibo_search_test.xlsx", [("可爱", "公园")], "test-key", client=client,
        )
        await reviewer.add_note({"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        assert sent == []
        await reviewer.finalize_note("1")
        assert len(sent) == 2 and "这个公园好可爱" in sent[0]
        assert reviewer.output_path.exists()
        await reviewer.add_note({"note_id": "2", "content": "公园里的小狗好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await reviewer.add_comment({"comment_id": "2a", "note_id": "2", "content": "这个公园好可爱"})
        await reviewer.add_comment({"comment_id": "2b", "note_id": "2", "content": "这里的树也很漂亮"})
        assert len(sent) == 2
        await reviewer.finalize_note("2")
        await reviewer.close()

    result = openpyxl.load_workbook(reviewer.output_path)
    assert result["Matched"].max_row == 3
    assert "Excluded" not in result.sheetnames
    assert result["MatchedComments"].max_row == 2
    assert dict(result["Summary"].values)["api_calls"] == 4
    assert "一级评论 2b" in sent[2]


def test_long_note_context_sends_every_comment_without_silent_truncation():
    state = LiveNote(
        note={"note_id": "1", "title": "这个公园", "content": "公园介绍", "source_keyword": "可爱 公园", "image_names": "001.jpg"},
        pair=("可爱", "公园"),
        comments=[
            {"comment_id": str(index), "content": f"评论{index}:" + "很可爱" * 150}
            for index in range(12)
        ],
    )
    chunks = _note_context_chunks(state)
    assert len(chunks) > 1
    assert all(len(text) <= 3000 for text, _ in chunks)
    combined = "".join(text for text, _ in chunks)
    assert all(f"评论{index}:" in combined for index in range(12))
    assert all(f"一级评论 {index}" in combined for index in range(12))


@pytest.mark.asyncio
async def test_completed_note_resumes_from_checkpoint_without_new_model_call(tmp_path):
    calls = []

    def respond(request):
        calls.append(request)
        return grounded_response(request)

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        first = StreamingSemanticFilter(
            tmp_path / "weibo_search_first.xlsx", [("可爱", "公园")], "test-key", client=client,
        )
        await first.add_note({"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await first.add_comment({"note_id": "1", "comment_id": "a", "content": "真好看"})
        await first.finalize_note("1")
        await first.close()
        assert len(calls) == 2
        resumed = StreamingSemanticFilter(
            tmp_path / "weibo_search_second.xlsx", [("可爱", "公园")], "test-key", client=client,
        )
        await resumed.add_note({"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
        await resumed.finalize_note("1")
        await resumed.close()
    assert len(calls) == 2
    assert resumed.states[("1", "可爱 公园")].status == "matched"
    assert resumed.output_path.exists()


@pytest.mark.asyncio
async def test_excel_resume_restores_collected_comments(tmp_path, monkeypatch):
    from store.excel_store_base import ExcelStoreBase
    import config

    def respond(request):
        return grounded_response(request)

    data_dir = tmp_path / "weibo"
    data_dir.mkdir()
    note = {"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"}
    comment = {"note_id": "1", "comment_id": "a", "content": "好可爱"}
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        first = StreamingSemanticFilter(data_dir / "first.xlsx", [("可爱", "公园")], "test-key", client=client, comment_budget=config.CRAWLER_MAX_COMMENTS_COUNT_SINGLENOTES)
        await first.add_note(note)
        await first.add_comment(comment)
        await first.finalize_note("1")
        await first.close()

    monkeypatch.setattr(config, "SAVE_DATA_PATH", str(tmp_path))
    monkeypatch.setenv("MEDIACRAWLER_SEMANTIC_FILTER", "1")
    monkeypatch.setenv("MEDIACRAWLER_KEYWORD_PAIRS", '[["可爱", "公园"]]')
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-key")
    monkeypatch.setattr(ExcelStoreBase, "_instances", {})
    instance = ExcelStoreBase.get_instance("weibo", "search")
    await instance.store_content(note)
    assert ExcelStoreBase.semantic_note_completed("weibo", "1", "可爱 公园")
    await ExcelStoreBase.finalize_semantic_note("weibo", "1")
    await ExcelStoreBase.finalize_semantics_all()
    instance.flush()
    restored = openpyxl.load_workbook(instance.filename)
    assert list(restored["Contents"].values)[1][0] == "1"
    assert list(restored["Comments"].values)[1][1] == "a"


@pytest.mark.asyncio
async def test_strict_mode_preserves_uncertainty_but_not_reject_with_missing_comments(tmp_path):
    def respond(request):
        text = json.loads(request.content)["state"]["text"]
        choice = (
            "reject" if "小狗" in text else
            "review" if "也许" in text else "keep"
        )
        return grounded_response(request, choice, relation='unclear' if choice == 'review' else 'target_itself')

    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        reviewer = StreamingSemanticFilter(
            tmp_path / "weibo_search_test.xlsx", [("可爱", "公园")], "test-key", client=client, selection_mode='strict',
        )
        for note_id, content in (
            ("1", "公园里的小狗好可爱"),
            ("2", "这个公园也许很可爱"),
            ("3", "这个公园好可爱"),
        ):
            await reviewer.add_note({"note_id": note_id, "content": content, "source_keyword": "可爱 公园", "image_names": "001.jpg"})
            await reviewer.finalize_note(note_id, comments_available=False)
        await reviewer.add_note({
            "note_id": "4", "content": "公园里的小狗好可爱",
            "source_keyword": "可爱 公园", "semantic_full_text_available": False,
                                               "image_names": "001.jpg",
        })
        await reviewer.finalize_note("4")
        await reviewer.close()

    result = openpyxl.load_workbook(reviewer.output_path)
    assert reviewer.states[("1", "可爱 公园")].status == "excluded"
    assert reviewer.states[("1", "可爱 公园")].source == "semantic_reject"
    assert [row[0] for row in list(result["Review"].values)[1:]] == ["2"]
    assert reviewer.states[("4", "可爱 公园")].source == "semantic_reject"
    assert [row[0] for row in list(result["Matched"].values)[1:]] == ["3"]
    assert dict(result["Summary"].values)["excluded"] == 2


def test_retained_count_is_specific_to_current_keyword_pair(monkeypatch):
    from store.excel_store_base import ExcelStoreBase

    stream = SimpleNamespace(
        states={
            ("1", "可爱 公园"): SimpleNamespace(status="matched"),
            ("1", "可爱 广场"): SimpleNamespace(status="excluded"),
        },
        note_keys={"1": {("1", "可爱 公园"), ("1", "可爱 广场")}},
    )
    monkeypatch.setattr(
        ExcelStoreBase, "_instances", {"weibo_search": SimpleNamespace(semantic_stream=stream)},
    )

    assert ExcelStoreBase.semantic_note_retained("weibo", "1", "可爱 公园")
    assert not ExcelStoreBase.semantic_note_retained("weibo", "1", "可爱 广场")


@pytest.mark.asyncio
async def test_excel_store_passes_live_records_to_semantic_filter(monkeypatch, tmp_path):
    import main as crawler_main
    import semantic_filter
    from store.excel_store_base import ExcelStoreBase

    monkeypatch.setattr(crawler_main.config, "SAVE_DATA_PATH", str(tmp_path))
    monkeypatch.setattr(crawler_main.config, "SAVE_DATA_OPTION", "excel")
    monkeypatch.setattr(crawler_main.config, "CRAWLER_TYPE", "search")
    monkeypatch.setattr(crawler_main.config, "PLATFORM", "wb")
    monkeypatch.setenv("MEDIACRAWLER_SEMANTIC_FILTER", "1")
    monkeypatch.setenv("MEDIACRAWLER_KEYWORD_PAIRS", '[["可爱", "公园"]]')
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-key")
    monkeypatch.setattr(ExcelStoreBase, "_instances", {})
    monkeypatch.setattr(crawler_main, "_excel_flushed", False)
    events = []

    class FakeStream:
        def __init__(self, raw_path, pairs, key, model, comment_budget=None):
            events.append(("created", raw_path, pairs, key, model))
            self.note_keys = {}
            self.states = {}
        async def add_note(self, note):
            events.append(("note", note["note_id"]))
            key = (note["note_id"], note["source_keyword"])
            self.note_keys.setdefault(note["note_id"], set()).add(key)
            self.states[key] = SimpleNamespace(status=None, note=note)
        async def add_comment(self, comment):
            events.append(("comment", comment["note_id"]))
        async def finalize_note(self, note_id, comments_available):
            events.append(("finalize", note_id, comments_available))
            for key in self.note_keys[note_id]:
                self.states[key].status = {"1": "matched", "2": "excluded", "3": "review"}[note_id]
        async def close(self):
            events.append(("close",))

    monkeypatch.setattr(semantic_filter, "StreamingSemanticFilter", FakeStream)
    instance = ExcelStoreBase.get_instance("weibo", "search")
    await instance.store_content({"note_id": "1", "content": "这个公园好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
    await instance.store_comment({"comment_id": "a", "note_id": "1", "content": "确实如此"})
    await ExcelStoreBase.finalize_semantic_note("weibo", "1")
    await instance.store_content({"note_id": "2", "content": "公园里的小狗好可爱", "source_keyword": "可爱 公园", "image_names": "001.jpg"})
    await instance.store_comment({"comment_id": "b", "note_id": "2", "content": "确实如此"})
    await ExcelStoreBase.finalize_semantic_note("weibo", "2")
    await instance.store_content({"note_id": "3", "content": "这个公园也许很可爱", "source_keyword": "可爱 公园", "comments_count": "0", "image_names": "001.jpg"})
    await instance.store_comment({"comment_id": "c", "note_id": "3", "content": "不确定"})
    await ExcelStoreBase.finalize_semantic_note("weibo", "3")
    assert ExcelStoreBase.semantic_note_retained("weibo", "1")
    assert not ExcelStoreBase.semantic_note_retained("weibo", "2")
    assert not ExcelStoreBase.semantic_note_retained("weibo", "3")
    await crawler_main._finalize_semantics_if_needed()
    crawler_main._flush_excel_if_needed()
    assert [event[0] for event in events] == ["created", "note", "comment", "finalize", "note", "comment", "finalize", "note", "comment", "finalize", "close"]
    assert instance.filename.exists()
    raw = openpyxl.load_workbook(instance.filename)
    assert [row[0] for row in list(raw["Contents"].values)[1:]] == ["1"]
    assert [row[0] for row in list(raw["Review"].values)[1:]] == ["3"]
    assert list(raw["Contents"].values)[0][:4] == (
        "note_id", "semantic_result", "semantic_source", "semantic_confidence",
    )
    assert [row[1] for row in list(raw["Contents"].values)[1:]] == ["matched"]
    assert [row[0] for row in list(raw["Comments"].values)[1:]] == ["a"]
    assert [row[0] for row in list(raw["ReviewComments"].values)[1:]] == ["c"]


@pytest.mark.asyncio
async def test_weibo_empty_comment_response_is_normal_and_not_retried(monkeypatch):
    from media_platform.weibo.client import WeiboClient
    import media_platform.weibo.client as weibo_client_module

    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json={"ok": 0, "msg": "还没有人评论哦~快来抢沙发！"})

    monkeypatch.setattr(
        weibo_client_module,
        "make_async_client",
        lambda proxy: httpx.AsyncClient(transport=httpx.MockTransport(respond)),
    )
    client = WeiboClient.__new__(WeiboClient)
    client.proxy = None
    client.timeout = 5
    client._host = "https://m.weibo.cn"
    client.headers = {}

    async def no_refresh():
        pass

    client._refresh_proxy_if_expired = no_refresh
    comments = await client.get_note_all_comments("note-1", crawl_interval=0)
    assert comments == []
    assert len(requests) == 1


@pytest.mark.asyncio
async def test_weibo_collects_all_top_level_comment_pages_without_subcomments():
    from media_platform.weibo.client import WeiboClient

    client = WeiboClient.__new__(WeiboClient)
    cursors = []
    callbacks = []

    async def page(note_id, max_id, max_id_type):
        cursors.append(max_id)
        if max_id == -1:
            return {"max_id": 2, "max_id_type": 0, "data": [{"id": str(i)} for i in range(10)]}
        return {"max_id": 0, "max_id_type": 0, "data": [{"id": str(i)} for i in range(10, 13)]}

    async def callback(note_id, items):
        callbacks.extend(items)

    async def no_subcomments(*args, **kwargs):
        raise AssertionError("subcomments must not be requested")

    client.get_note_comments = page
    client.get_comments_all_sub_comments = no_subcomments
    comments = await client.get_note_all_comments(
        "note-1", crawl_interval=0, callback=callback,
        max_count=None, include_sub_comments=False,
    )
    assert cursors == [-1, 2]
    assert len(comments) == len(callbacks) == 13
