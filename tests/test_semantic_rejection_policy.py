"""Definite Jev negatives win over incomplete fetches and review promotion."""
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import openpyxl
import pytest

from semantic_filter import Decision, JevReviewer, StreamingSemanticFilter
from semantic_quality import explicit_jev_rejection
from tests.semantic_responses import grounded_response


@pytest.mark.asyncio
@pytest.mark.parametrize('mode', ['candidate', 'strict'])
@pytest.mark.parametrize('choice,kind,reason', [
    ('reject', 'place', 'semantic_reject'),
    ('reject', 'animal', 'semantic_reject'),
    ('reject', 'pet', 'pet_not_allowed'),
    ('review', 'pet', 'pet_not_allowed'),
    ('keep', 'pet', 'pet_not_allowed'),
])
@pytest.mark.parametrize('stop', ['comment_budget_exhausted', 'comment_page_budget_exhausted', 'comments_unavailable'])
async def test_explicit_negative_is_excluded_even_with_incomplete_comments(tmp_path, mode, choice, kind, reason, stop):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: grounded_response(
        request, choice, kind=kind, relation='scene_component', confidence=None))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client,selection_mode=mode)
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'带我家可爱的小狗去公园','image_names':'001.jpg','semantic_full_text_available':False})
        await stream.finalize_note('n',False,interim=True)
        state=stream.states[('n','可爱 公园')]
        assert state.status=='excluded' and state.source==reason
        assert not state.matched and not state.uncertain
        assert state.note['jev_original_choice']==choice
        assert state.note['semantic_rejection_reason']==reason
        assert 'selection_origin' not in state.note
        await stream.add_comment({'note_id':'n','comment_id':'later','content':'公园的花坛好可爱'})
        await stream.finalize_note('n',False,comment_stop_reason=stop)
        assert not state.comments and state.status=='excluded'
        assert stream.reviewer.calls==1
        checkpoint=stream.checkpoint_path
        await stream.close()
        resumed=StreamingSemanticFilter(tmp_path/'resumed.xlsx',[('可爱','公园')],'fake',client=client,selection_mode=mode)
        assert resumed.checkpoint_path==checkpoint
        await resumed.add_note(dict(state.note))
        await resumed.finalize_note('n',False,interim=True)
        assert resumed.states[('n','可爱 公园')].rejection_reason==reason
        assert resumed.reviewer.calls==0
        await resumed.close()
    book=openpyxl.load_workbook(stream.output_path,read_only=True)
    assert dict(book['Summary'].values)['excluded']==1
    assert book['Matched'].max_row==book['Review'].max_row==1
    book.close()


@pytest.mark.asyncio
async def test_later_rejection_overrides_earlier_uncertainty(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: grounded_response(
        request,'reject' if 'comment:late' in json.dumps(json.loads(request.content)) else 'review',kind='unclear',relation='unclear'))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client,selection_mode='strict')
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'公园也许可爱','image_names':'001.jpg'})
        await stream.finalize_note('n',False,interim=True)
        state=stream.states[('n','可爱 公园')]
        assert state.uncertain and state.status is None
        await stream.add_comment({'note_id':'n','comment_id':'late','content':'公园这个名字可爱'})
        await stream.finalize_note('n',False,interim=True)
        assert state.status=='excluded' and state.rejection_reason=='semantic_reject'
        await stream.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('mode', ['candidate','strict'])
@pytest.mark.parametrize('support,expected', [('contradicted','reject'),('unclear','review')])
async def test_verification_negative_is_distinct_from_uncertainty(mode, support, expected):
    def respond(request):
        if 'support' in json.loads(request.content)['questions']:
            return httpx.Response(200,json={'answers':{'support':{'choice':support,'confidence':.9}}})
        return grounded_response(request)
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        decision=await JevReviewer('fake',client,selection_mode=mode).decide(('可爱','公园'),'公园的花坛好可爱')
    # Candidate mode may retain grounded unclear evidence, never contradicted.
    assert decision.choice==(expected if mode=='strict' or support=='contradicted' else 'keep')


@pytest.mark.asyncio
async def test_stream_cannot_promote_a_transformed_original_reject(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request:pytest.fail('Replay only'))) as client:
        stream=StreamingSemanticFilter(tmp_path/'run.xlsx',[('可爱','公园')],'fake',client=client)
        stream.reviewer.decide=AsyncMock(return_value=Decision('review',object_type='animal',original_model_choice='reject'))
        await stream.add_note({'note_id':'n','source_keyword':'可爱 公园','content':'公园的小动物好可爱','image_names':'001.jpg'})
        await stream.finalize_note('n',False,comment_stop_reason='comment_budget_exhausted')
        assert stream.states[('n','可爱 公园')].status=='excluded'
        await stream.close()


@pytest.mark.asyncio
async def test_rejected_posts_do_not_use_quota_comments_downloads_or_exports(tmp_path,monkeypatch):
    import config
    from media_platform.weibo import core
    from store.excel_store_base import ExcelStoreBase
    for name,value in {'KEYWORDS':'可爱 公园','START_PAGE':1,'CRAWLER_MAX_NOTES_COUNT':1,'CRAWLER_MAX_SLEEP_SEC':0,
                       'ENABLE_GET_COMMENTS':True,'SAVE_DATA_OPTION':'excel','SAVE_DATA_PATH':str(tmp_path)}.items():
        monkeypatch.setattr(config,name,value)
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER',raising=False)
    store=ExcelStoreBase('weibo','search')
    monkeypatch.setenv('MEDIACRAWLER_SEMANTIC_FILTER','1')
    monkeypatch.setattr(ExcelStoreBase,'_instances',{'weibo_search':store})
    monkeypatch.setattr(core.weibo_store.WeibostoreFactory,'create_store',lambda:store)
    monkeypatch.setattr(core,'filter_search_result_card',lambda cards:cards)
    cards=[{'mblog':{'id':nid,'text':text,'created_at':'Thu Oct 08 09:00:00 +0800 2026',
                     'pics':[{'url':'https://example.com/test.jpg'}]}}
           for nid,text in [('pet','带我家可爱的小狗去公园'),('rejected','公园这个名字好可爱'),('kept','公园里的花坛好可爱')]]
    crawler=core.WeiboCrawler()
    crawler.wb_client=SimpleNamespace(get_note_by_keyword=AsyncMock(return_value={'cards':cards}))
    crawler.batch_get_notes_full_text=AsyncMock(side_effect=lambda notes:notes)
    async def complete_comments(ids):
        assert ids == ['kept']  # Explicit rejects never fetch comments; accepted posts still deliver them.
        await ExcelStoreBase.finalize_semantic_note('weibo','kept',True,comment_stop_reason='comments_exhausted')
    crawler.batch_get_notes_comments=AsyncMock(side_effect=complete_comments)
    crawler.download_media=AsyncMock()
    def respond(request):
        text=json.loads(request.content)['state']['text']
        return grounded_response(request,'review' if '小狗' in text else 'reject' if '名字' in text else 'keep',kind='pet' if '小狗' in text else 'place')
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        store.semantic_stream=StreamingSemanticFilter(store.filename,[('可爱','公园')],'fake',client=client)
        await crawler.search()
        assert [call.args[0]['id'] for call in crawler.download_media.await_args_list]==['kept']
        assert store.semantic_stream.states[('pet','可爱 公园')].source=='pet_not_allowed'
        assert store.semantic_stream.states[('rejected','可爱 公园')].source=='semantic_reject'
        await store.semantic_stream.close()
    store.flush()
    book=openpyxl.load_workbook(store.filename,read_only=True)
    values=iter(book['Contents'].values); headers=next(values)
    assert [dict(zip(headers,row))['note_id'] for row in values]==['kept']
    book.close()


def test_only_definite_negatives_trigger_hard_gate():
    assert explicit_jev_rejection('review','animal') is None
    assert explicit_jev_rejection('review','unclear') is None
    assert explicit_jev_rejection('keep','scene_animal') is None
    assert explicit_jev_rejection('keep','park_animal') is None
