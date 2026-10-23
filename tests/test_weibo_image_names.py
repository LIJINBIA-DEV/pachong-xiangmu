"""Filename-only Weibo metadata through both Excel exports and local browsing."""
import json
from unittest.mock import AsyncMock

import httpx
import openpyxl
import pytest

from api.services.data_browser import browse_posts
from semantic_filter import LiveNote, StreamingSemanticFilter
from store import weibo
from store.excel_store_base import ExcelStoreBase
from tests.semantic_responses import grounded_response
from var import source_keyword_var


@pytest.mark.asyncio
@pytest.mark.parametrize(('media', 'expected'), [
    ({'pics': [{'url': 'https://wx4.sinaimg.cn/orj480/a.jpg'}, {'url': 'https://wx4.sinaimg.cn/orj480/b.png'}]}, '001.jpg,002.png'),
    ({'retweeted_status': {'pics': [{'url': 'https://wx4.sinaimg.cn/orj480/a.jpg'}]}}, '001.jpg'),
    ({'page_info': {'type': 'video', 'media_info': {'stream_url': 'https://example.com/video.mp4'}, 'page_pic': {'url': 'https://example.com/cover.webp'}}}, 'cover.webp'),
    ({}, ''),
])
async def test_weibo_record_stores_names_without_image_urls(monkeypatch, media, expected):
    captured = AsyncMock()
    monkeypatch.setattr(weibo.WeibostoreFactory, 'create_store', lambda: captured)
    await weibo.update_weibo_note({'mblog': {
        'id': '123', 'text': '公园的花坛真可爱',
        'created_at': 'Thu Oct 08 09:00:00 +0800 2026', **media,
    }})
    note = captured.store_content.call_args.kwargs['content_item']
    assert note['image_names'] == expected
    assert 'image_list' not in note
    assert 'sinaimg' not in json.dumps(note)
    assert 'example.com' not in json.dumps(note)
    assert note['note_url'] == 'https://m.weibo.cn/detail/123'


@pytest.mark.asyncio
async def test_filename_metadata_survives_primary_and_semantic_exports_and_browser(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER', raising=False)
    store = ExcelStoreBase('weibo', 'search')
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(store.filename, [('可爱', '公园')], 'fake', client=client)
        store.semantic_stream = stream
        monkeypatch.setattr(weibo.WeibostoreFactory, 'create_store', lambda: store)
        token = source_keyword_var.set('可爱 公园')
        try:
            await weibo.update_weibo_note({'mblog': {
                'id': '123', 'text': '这个公园的花坛真可爱',
                'user': {'id': 'author-123', 'screen_name': '公园观察员'},
                'created_at': 'Thu Oct 08 09:00:00 +0800 2026',
                'pics': [{'url': 'https://wx4.sinaimg.cn/orj480/a.jpg'}],
            }})
        finally:
            source_keyword_var.reset(token)
        await stream.finalize_note('123')
        store._write_retained_semantic_note('123')
        await stream.close()
        store.flush()
    for path, sheet in [(store.filename, 'Contents'), (stream.output_path, 'Matched')]:
        book = openpyxl.load_workbook(path, read_only=True)
        try:
            headers, row = list(book[sheet].values)
            assert dict(zip(headers, row))['image_names'] == '001.jpg'
            assert dict(zip(headers, row))['nickname'] == '公园观察员'
            assert 'image_list' not in headers
        finally:
            book.close()
    checkpoint = [json.loads(line) for line in stream.checkpoint_path.read_text(encoding='utf-8').splitlines()]
    assert checkpoint[-1]['note']['image_names'] == '001.jpg'
    assert 'sinaimg' not in json.dumps(checkpoint)
    media_dir = tmp_path / 'wb' / 'media' / '123'
    media_dir.mkdir(parents=True)
    (media_dir / '001.jpg').write_bytes(b'local-image-fixture')
    for path in (store.filename, stream.output_path):
        post = browse_posts(tmp_path, path)['posts'][0]
        assert post['nickname'] == '公园观察员'
        assert post['images'] == [{'url': '/api/data/media/wb/media/123/001.jpg', 'name': '001.jpg', 'local': True}]
        assert post['status'] == 'matched'


@pytest.mark.asyncio
@pytest.mark.parametrize('status', ['matched', 'review', 'pending', 'duplicate'])
async def test_legacy_container_fields_no_longer_crash_semantic_excel(tmp_path, status):
    note = {'note_id': 'legacy', 'source_keyword': '可爱 公园',
            'image_list': ['https://example.com/old.jpg'], 'metadata': {'说明': '猫形花坛'},
            'likes': 12, 'available': True}
    comment = {'comment_id': 'c1', 'note_id': 'legacy', 'content': '公园好可爱', 'tags': ['花坛', '猫']}
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        stream = StreamingSemanticFilter(tmp_path / 'legacy.xlsx', [('可爱', '公园')], 'fake', client=client)
        state = LiveNote(note=note, pair=('可爱', '公园'), status=status, matched_comment=comment)
        key = ('legacy', '可爱 公园')
        stream.states[key] = state
        stream.active_keys.add(key)
        stream.comment_headers = list(comment)
        stream.save()
        await stream.close()
    book = openpyxl.load_workbook(stream.output_path, read_only=True)
    try:
        sheet = {'matched': 'Matched', 'review': 'Review', 'pending': 'Pending', 'duplicate': 'Duplicates'}[status]
        headers, row = list(book[sheet].values)
        record = dict(zip(headers, row))
        assert json.loads(record['image_list']) == note['image_list']
        assert json.loads(record['metadata']) == note['metadata']
        assert record['likes'] == 12 and record['available'] is True
        if status == 'matched':
            headers, row = list(book['MatchedComments'].values)
            assert json.loads(dict(zip(headers, row))['tags']) == comment['tags']
    finally:
        book.close()
    assert isinstance(note['image_list'], list)  # Checkpoint evidence stays structured.


@pytest.mark.asyncio
async def test_cached_decision_refreshes_media_names_without_rejudging(tmp_path):
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        old = StreamingSemanticFilter(tmp_path / 'old.xlsx', [('可爱', '公园')], 'fake', client=client)
        await old.add_note({'note_id': '1', 'nickname': '公***员', 'source_keyword': '可爱 公园', 'content': '这个公园好可爱', 'image_list': ['https://example.com/old.jpg']})
        await old.finalize_note('1')
        await old.close()
        resumed = StreamingSemanticFilter(tmp_path / 'new.xlsx', [('可爱', '公园')], 'fake', client=client)
        await resumed.add_note({'note_id': '1', 'nickname': '公园观察员', 'source_keyword': '可爱 公园', 'content': '这个公园好可爱', 'image_names': '001.jpg'})
        await resumed.close()
        assert resumed.reviewer.calls == 0
        assert resumed.states[('1', '可爱 公园')].status == 'matched'
    book = openpyxl.load_workbook(resumed.output_path, read_only=True)
    try:
        headers, row = list(book['Matched'].values)
        assert dict(zip(headers, row))['image_names'] == '001.jpg'
        assert dict(zip(headers, row))['nickname'] == '公园观察员'
        assert 'image_list' not in headers
    finally:
        book.close()
