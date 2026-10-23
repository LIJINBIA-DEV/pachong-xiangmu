"""Public metadata survives collection, browsing, review and delivery."""
import copy
import asyncio
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import config
from api.routers import data as router
from api.services.manual_review import update_annotation
from store.excel_store_base import ExcelStoreBase
from store.weibo import update_weibo_note
from var import crawler_type_var, source_keyword_var


def note(nid='1234567890', user=None, **fields):
    return {'mblog': {'id': nid, 'text': '公园里的雕塑真可爱',
        'created_at': 'Sat Jun 14 12:00:00 +0800 2025',
        'user': user if user is not None else {'idstr': '765432198765432198',
            'screen_name': '公园散步的人', 'verified': False, 'gender': 'f',
            'profile_url': 'https://m.weibo.cn/profile/765432198765432198'}, **fields}}


@pytest.fixture
def collection(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'SAVE_DATA_PATH', str(tmp_path))
    monkeypatch.setattr(config, 'SAVE_DATA_OPTION', 'excel')
    monkeypatch.setattr(config, 'CRAWLER_TYPE', 'search')
    monkeypatch.delenv('MEDIACRAWLER_SEMANTIC_FILTER', raising=False)
    monkeypatch.setattr(config, 'CRAWLER_MAX_SLEEP_SEC', 0)
    ExcelStoreBase._instances.clear()
    ct = crawler_type_var.set('search')
    kw = source_keyword_var.set('可爱 公园')
    app = FastAPI()
    app.include_router(router.router, prefix='/api')
    yield tmp_path, TestClient(app)
    ExcelStoreBase._instances.clear()
    crawler_type_var.reset(ct)
    source_keyword_var.reset(kw)


async def collect(root, client, notes):
    for item in notes:
        await update_weibo_note(copy.deepcopy(item))
    store = ExcelStoreBase.get_instance('weibo', 'search')
    store.flush()
    path = Path(store.filename)
    file = path.relative_to(root).as_posix()
    response = client.get('/api/data/browse', params={'file': file})
    assert response.status_code == 200, response.text
    posts = response.json()['posts']
    for post in posts:
        update_annotation(root, path, post['key'], 'post', 'qualified')
    before = path.read_bytes()
    response = client.post('/api/data/review/export', json={'file': file})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    headers, *rows = list(book['审核结果'].values)
    exported = [dict(zip(headers, row)) for row in rows]
    book.close()
    assert path.read_bytes() == before
    return posts, exported, path


@pytest.mark.asyncio
async def test_public_account_and_gender_survive_collection_and_delivery(collection):
    item = note()
    item['mblog']['user']['profile_url'] += '?user_token=visitor-token#tracking'
    posts, rows, path = await collect(*collection, [item])
    assert posts[0]['account_id'] == '765432198765432198'
    assert posts[0]['gender'] == '女'
    assert posts[0]['profile_url'] == 'https://m.weibo.cn/profile/765432198765432198'
    assert posts[0]['verification'] == '无认证说明'
    assert rows[0]['账号昵称'] == '公园散步的人'
    assert rows[0]['账号 ID'] == '765432198765432198'
    assert rows[0]['性别'] == '女'
    assert rows[0]['账号主页链接'] == 'https://m.weibo.cn/profile/765432198765432198'
    assert len(rows[0]) == 16
    raw = load_workbook(path)
    records = list(raw['Contents'].values)
    assert records[1][records[0].index('creator_hash')] != '765432198765432198'
    raw.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('fields,expected', [
    ({'annotations': [{'place': {'title': '南京·玄武湖公园', 'poiid': 'B209'}}]}, '南京·玄武湖公园'),
    ({'geo': {'type': 'Point', 'coordinates': [32.07, 118.8]}}, '经度 118.8，纬度 32.07'),
    ({'geo': {'type': 'Point', 'coordinates': [192, 118.8]}}, '未获取'),
    ({'geo': {}, 'annotations': [], 'location': '账号里含义不明的地点'}, '未获取'),
    ({'retweeted_status': {'annotations': [{'place': {'title': '别人的定位'}}]}}, '未获取'),
])
async def test_only_explicit_post_location_reaches_delivery(collection, fields, expected):
    posts, rows, _ = await collect(*collection, [note(**fields)])
    assert posts[0]['display_location'] == rows[0]['地点（定位优先）'] == expected
    source = '未获取' if expected == '未获取' else '微博定位'
    assert posts[0]['location_source'] == rows[0]['地点来源'] == source


@pytest.mark.asyncio
async def test_missing_and_unknown_profiles_keep_stable_excel_columns(collection):
    missing = note('1', user={'screen_name': '资料缺失'})
    unknown = note('2', user={'idstr': '7654321', 'screen_name': '未知性别',
        'gender': 'n', 'profile_url': 'https://m.weibo.cn.evil.example/u/1', 'verified': False})
    complete = note('3', user={'idstr': '765432198765432198', 'screen_name': '完整资料',
        'gender': 'm', 'verified': True, 'verified_reason': '=摄影师',
        'profile_url': 'javascript:alert(1)'})
    posts, rows, path = await collect(*collection, [missing, unknown, complete])
    assert [p['gender'] for p in posts] == ['未获取', '未知', '男']
    assert [r['账号 ID'] for r in rows] == ['未获取', '7654321', '765432198765432198']
    assert [p['profile_url'] for p in posts] == ['', '', '']
    assert [r['认证说明'] for r in rows] == ['未获取', '无认证说明', '=摄影师']
    raw = load_workbook(path)
    data = list(raw['Contents'].values)
    col = data[0].index('author_verification') + 1
    assert raw['Contents'].cell(4, col).data_type == 's'
    raw.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('fields,expected,source,region', [
    ({'region_name': '发布于 江苏', 'annotations': [{'place': {'title': '南京·玄武湖'}}]}, '南京·玄武湖', '微博定位', '江苏'),
    ({'region_name': 'IP属地：上海'}, '上海', 'IP 属地', '上海'),
    ({'region_name': '发布于 海外'}, '海外', 'IP 属地', '海外'),
    ({'region_name': 'IP属地：'}, '未获取', '未获取', ''),
    ({'region_name': {'name': '不支持的对象'}}, '未获取', '未获取', ''),
    ({'region_name': '127.0.0.1'}, '未获取', '未获取', ''),
    ({'user': {'screen_name': '散步的人', 'ip_location': '广东', 'location': '广州'}}, '未获取', '未获取', ''),
])
async def test_post_ip_fallback_has_explicit_source(collection, fields, expected, source, region):
    posts, rows, _ = await collect(*collection, [note(**fields)])
    assert posts[0]['display_location'] == rows[0]['地点（定位优先）'] == expected
    assert posts[0]['location_source'] == rows[0]['地点来源'] == source
    assert posts[0]['post_ip_region'] == region
    assert rows[0]['IP 属地'] == (region or '未获取')


@pytest.mark.asyncio
async def test_new_collection_enriches_once_per_author_and_preserves_post_location(collection):
    from media_platform.weibo.core import WeiboCrawler

    class Platform:
        def __init__(self):
            self.profiles = []
            self.details = []

        async def get_creator_info_by_id(self, creator_id):
            self.profiles.append(creator_id)
            await asyncio.sleep(0)
            return {'userInfo': {'idstr': creator_id, 'gender': 'f', 'verified': False,
                'profile_url': 'https://m.weibo.cn/profile/' + creator_id}}

        async def get_note_info_by_id(self, nid):
            self.details.append(nid)
            return note(nid, text='短文', region_name='发布于 江苏' if nid == '1' else '发布于 上海')

    platform = Platform()
    crawler = WeiboCrawler()
    crawler.wb_client = platform
    original = [note(str(i), user={'idstr': '7654321', 'screen_name': '散步的人'}, text='完整原文，不能被较短详情覆盖') for i in (1, 2)]
    prepared = await asyncio.gather(*(crawler.prepare_note(item) for item in original))
    posts, rows, _ = await collect(*collection, prepared)
    assert platform.profiles == ['7654321']
    assert sorted(platform.details) == ['1', '2']
    assert [p['gender'] for p in posts] == ['女', '女']
    assert [r['地点（定位优先）'] for r in rows] == ['江苏', '上海']
    assert all(p['content'] == '完整原文，不能被较短详情覆盖' for p in posts)
    assert all(p['account_id'] == '7654321' for p in posts)


@pytest.mark.asyncio
@pytest.mark.parametrize('response', [None, {'data': 'not an object'}, {'userInfo': []}])
async def test_failed_or_malformed_metadata_does_not_drop_posts(collection, response):
    from media_platform.weibo.core import WeiboCrawler

    class Platform:
        async def get_note_info_by_id(self, nid):
            raise RuntimeError('network unavailable')

        async def get_creator_info_by_id(self, uid):
            if response is None:
                raise RuntimeError('network unavailable')
            return response

    crawler = WeiboCrawler()
    crawler.wb_client = Platform()
    item = await crawler.prepare_note(note(user={'id': '7654321', 'screen_name': '散步的人'}))
    posts, rows, _ = await collect(*collection, [item])
    assert len(posts) == len(rows) == 1
    assert rows[0]['性别'] == rows[0]['地点（定位优先）'] == '未获取'


@pytest.mark.asyncio
async def test_resumed_semantic_result_keeps_new_metadata_without_reclassifying(collection):
    import httpx
    from semantic_filter import StreamingSemanticFilter
    from tests.semantic_responses import grounded_response

    root, api = collection
    calls = []
    def respond(request):
        calls.append(request)
        return grounded_response(request)
    raw = root / 'weibo' / 'resume.xlsx'
    raw.parent.mkdir(exist_ok=True)
    first = {'note_id': '1', 'content': '这个公园好可爱', 'source_keyword': '可爱 公园', 'image_names': '001.jpg'}
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        stream = StreamingSemanticFilter(raw, [('可爱', '公园')], 'test-key', client=client)
        await stream.add_note(first)
        await stream.finalize_note('1')
        await stream.close()
        before = len(calls)
        resumed = StreamingSemanticFilter(raw, [('可爱', '公园')], 'test-key', client=client)
        await resumed.add_note({**first, 'metadata_version': 1, 'account_id': '7654321',
            'gender': '女', 'publish_location': '玄武湖公园', 'post_ip_region': '江苏',
            'account_verified': True, 'account_verification': '公开认证说明'})
        await resumed.finalize_note('1')
        await resumed.close()
        restarted = StreamingSemanticFilter(root / 'weibo' / 'restarted.xlsx', [('可爱', '公园')], 'test-key', client=client)
        await restarted.add_note(first)
        await restarted.finalize_note('1')
        await restarted.close()
    assert len(calls) == before
    for path in (resumed.output_path, restarted.output_path):
        result = api.get('/api/data/browse', params={'file': path.relative_to(root).as_posix()})
        assert result.status_code == 200, result.text
        post = result.json()['posts'][0]
        assert post['account_id'] == '7654321'
        assert post['display_location'] == '玄武湖公园'
        assert post['verification'] == '公开认证说明'


@pytest.mark.asyncio
async def test_full_verification_text_survives_delivery(collection):
    item = note()
    description = '公开认证说明' * 80
    item['mblog']['user'].update(verified=True, verified_reason=description)
    posts, rows, _ = await collect(*collection, [item])
    assert posts[0]['verification'] == rows[0]['认证说明'] == description


@pytest.mark.asyncio
@pytest.mark.parametrize('complete_first', [True, False])
async def test_metadata_reused_across_search_rows(collection, complete_first):
    from media_platform.weibo.core import WeiboCrawler
    class Platform:
        async def get_note_info_by_id(self, nid):
            return {}
        async def get_creator_info_by_id(self, uid):
            return {}
    crawler = WeiboCrawler()
    crawler.wb_client = Platform()
    complete = note(annotations=[{'place': {'title': '公园定位'}}], region_name='发布于 江苏')
    missing = note(user={'idstr': '765432198765432198'})
    first = await crawler.prepare_note(complete if complete_first else missing)
    await update_weibo_note(first)
    source_keyword_var.set('萌 公园')
    second = await crawler.prepare_note(missing if complete_first else complete)
    posts, rows, _ = await collect(*collection, [second])
    assert len(posts) == 2
    assert {p['source_keyword'] for p in posts} == {'可爱 公园', '萌 公园'}
    assert all(p['gender'] == '女' and p['display_location'] == '公园定位' for p in posts)
    assert all(r['IP 属地'] == '江苏' for r in rows)


def test_other_platform_explicit_location_is_not_weibo(collection):
    from openpyxl import Workbook
    root, api = collection
    path = root / 'xhs' / 'sample.xlsx'
    path.parent.mkdir()
    book = Workbook(); sheet = book.active; sheet.title = 'Contents'
    sheet.append(['note_id', 'platform', 'content', 'publish_location', 'post_ip_region'])
    sheet.append(['1', 'xhs', '正文', '上海公园', '江苏'])
    sheet.append(['2', 'xhs', '正文', '', '江苏'])
    book.save(path); book.close()
    response = api.get('/api/data/browse', params={'file': 'xhs/sample.xlsx'})
    assert response.status_code == 200
    first, second = response.json()['posts']
    assert first['display_location'] == '上海公园'
    assert first['location_source'] == '帖子定位'
    assert second['display_location'] == '未获取'
    for post in (first, second):
        update_annotation(root, path, post['key'], 'post', 'qualified')
    export = api.post('/api/data/review/export', json={'file': 'xhs/sample.xlsx'})
    book = load_workbook(BytesIO(export.content))
    headers, *rows = list(book['审核结果'].values)
    assert rows[0][headers.index('地点来源')] == '帖子定位'
    assert rows[1][headers.index('地点（定位优先）')] == '未获取'
    book.close()


@pytest.mark.asyncio
async def test_later_metadata_updates_retained_semantic_search_rows(collection):
    import httpx
    from semantic_filter import StreamingSemanticFilter
    from tests.semantic_responses import grounded_response
    root, api = collection
    store = ExcelStoreBase.get_instance('weibo', 'search')
    first = {'note_id': '1', 'content': '公园里的雕塑真可爱又萌', 'source_keyword': '可爱 公园',
             'image_names': '001.jpg', 'metadata_version': 1, 'gender': ''}
    async with httpx.AsyncClient(transport=httpx.MockTransport(grounded_response)) as client:
        store.semantic_stream = StreamingSemanticFilter(store.filename, [('可爱', '公园'), ('萌', '公园')], 'test', client=client)
        await store.store_content(first)
        await ExcelStoreBase.finalize_semantic_note('weibo', '1')
        await store.store_content({**first, 'source_keyword': '萌 公园', 'gender': '女',
                                   'publish_location': '公园定位', 'account_verification': '公开认证'})
        await ExcelStoreBase.finalize_semantic_note('weibo', '1')
        await store.semantic_stream.close()
        store.flush()
    for path in (Path(store.filename), store.semantic_stream.output_path):
        result = api.get('/api/data/browse', params={'file': path.relative_to(root).as_posix()})
        assert result.status_code == 200
        posts = result.json()['posts']
        assert len(posts) == 2
        assert all(p['gender'] == '女' and p['display_location'] == '公园定位' and p['verification'] == '公开认证' for p in posts)
