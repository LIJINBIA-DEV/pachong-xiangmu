"""Exercise the workbench pipeline in an isolated data directory.

Only external model responses are simulated; the semantic exporter, task catalog,
review persistence, progress, logs and downloaded documents use real code.
"""
import asyncio
import hashlib
from io import BytesIO
import json

from fastapi import FastAPI
import httpx
from openpyxl import load_workbook
import pytest

from api.routers import data as data_router, dashboard as dashboard_router
from api.services import crawler_manager, content_summary
from collection_tasks import create_task, prepare_run, SearchProgress, finish_run
from semantic_filter import StreamingSemanticFilter


@pytest.mark.asyncio
async def test_filter_progress_review_export_summary_and_durable_logs(tmp_path, monkeypatch):
    monkeypatch.setattr(data_router, 'data_root', lambda: tmp_path)
    monkeypatch.setattr(dashboard_router, 'data_root', lambda: tmp_path)
    task = create_task(tmp_path, {'platform': 'xhs', 'semantic_filter': True,
        'semantic_rule': 'topic', 'enable_comments': True}, {'苏州 房价': 2})
    run = prepare_run(tmp_path, task, task['targets'])
    monkeypatch.setenv('MEDIACRAWLER_COLLECTION_RUN', str(run))
    monkeypatch.setattr(crawler_manager, 'collection_run', run)
    monkeypatch.setattr(crawler_manager, 'status', 'running')
    monkeypatch.setattr(crawler_manager, 'stage', 'running')

    decisions = iter([('keep', .95), ('keep', .4), ('reject', .95)])
    def semantic_response(request):
        payload = json.loads(request.content)
        choice, confidence = next(decisions)
        evidence = next(key for key in payload['questions']['evidence_text']['criteria'] if key != 'unclear')
        return httpx.Response(200, json={'answers': {
            'relation': {'choice': choice, 'confidence': confidence},
            'evidence_text': {'choice': evidence}}})

    output = tmp_path / 'xhs' / 'flow.xlsx'
    output.parent.mkdir()
    async with httpx.AsyncClient(transport=httpx.MockTransport(semantic_response)) as model:
        stream = StreamingSemanticFilter(output, [('苏州 房价', '')], 'fixture-key', client=model, semantic_rule='topic')
        for number in range(1, 4):
            await stream.add_note({'note_id': str(number), 'title': '苏州住宅价格讨论',
                'desc': f'苏州住宅成交总价约{number}百万元。', 'source_keyword': '苏州 房价',
                'note_url': f'https://www.xiaohongshu.com/explore/{number}',
                'image_list': 'https://example.org/post.jpg'})
            await stream.finalize_note(str(number))
        await stream.close()
    SearchProgress().record('苏州 房价', 1, ['1', '2', '3'], result={'retained': 1,
        'target': 2, 'reason': 'search_exhausted', 'diagnostics': {'evaluated': 3, 'pages_checked': 1}})
    file = stream.output_path.relative_to(tmp_path).as_posix()
    digest = hashlib.sha256(stream.output_path.read_bytes()).hexdigest()

    def summary_response(request):
        payload = json.loads(request.content)
        assert request.headers['Authorization'] == 'Bearer fixture-summary-key'
        data = json.loads(payload['messages'][1]['content'])
        if 'sources' in data['posts'][0]:
            result = {'posts': [{'ref': post['ref'], 'summary': '讨论住宅价格。',
                'stance': 'neutral', 'topics': ['住宅价格'],
                'evidence': [{'source': 'body', 'quote': next(s['text'] for s in post['sources'] if s['id'] == 'body')}]}
                for post in data['posts']]}
        else:
            result = {'overview': '帖子讨论苏州住宅成交价格。', 'themes': [
                {'title': '住宅价格', 'summary': '讨论住宅成交价格。', 'refs': ['P1', 'P2']}]}
        return httpx.Response(200, json={'choices': [{'finish_reason': 'stop',
            'message': {'content': json.dumps(result, ensure_ascii=False)}}]})
    jobs = content_summary.SummaryJobs(client_factory=lambda: httpx.AsyncClient(transport=httpx.MockTransport(summary_response)))
    monkeypatch.setattr(content_summary, 'summary_jobs', jobs)
    monkeypatch.setattr(data_router, 'summary_jobs', jobs)
    app = FastAPI()
    app.include_router(data_router.router, prefix='/api')
    app.include_router(dashboard_router.router, prefix='/api')
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://local') as api:
        catalog = (await api.get('/api/data/tasks')).json()['tasks'][0]
        assert catalog['counts'] == {'matched': 1, 'review': 1}
        progress = (await api.get('/api/dashboard/progress')).json()
        assert progress['checked'] == 3 and progress['target'] == 2 and progress['percent'] == 50
        review = (await api.get('/api/data/review', params={'file': file})).json()
        assert review['entries']['2:苏州 房价']['review_images'][0]['remote_url']
        finish_run(run, 'completed')
        monkeypatch.setattr(crawler_manager, 'status', 'idle')
        approved = await api.patch('/api/data/review', json={'file': file,
            'post_key': '2:苏州 房价', 'kind': 'post', 'decision': 'qualified'})
        assert approved.status_code == 200
        assert (await api.get('/api/data/tasks')).json()['tasks'][0]['counts'] == {'matched': 2}
        export = await api.get('/api/data/export', params={'file': file, 'status': 'matched'})
        assert export.status_code == 200
        book = load_workbook(BytesIO(export.content))
        assert book['帖子'].max_row == 3
        book.close()
        generated = await api.post('/api/data/content-summary', json={'file': file, 'api_key': 'fixture-summary-key'})
        assert generated.status_code == 200
        await asyncio.gather(*list(jobs.tasks.values()))
        report = (await api.get('/api/data/content-summary', params={'file': file})).json()['report']
        assert report['status'] == 'ready' and report['total'] == 2
        downloaded = await api.get(f"/api/data/content-summary/{report['id']}/download", params={'file': file})
        assert downloaded.status_code == 200 and '住宅价格' in downloaded.text
        logs = (await api.get('/api/data/collection-logs', params={'file': file})).json()
        assert {'run_started', 'run_finished', 'review_post', 'summary_finished'} <= {event['kind'] for event in logs['timeline']}
        assert 'fixture-key' not in json.dumps(logs) and 'fixture-summary-key' not in downloaded.text
    assert hashlib.sha256(stream.output_path.read_bytes()).hexdigest() == digest
